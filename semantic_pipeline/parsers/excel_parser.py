"""
ExcelParser — Merged-Cell-Aware Row-Level Tabular Chunker
==========================================================
Processes .xlsx files sheet-by-sheet with full merged cell resolution,
then stringifies each row by mapping column headers to cell values.

Pipeline:
  1. openpyxl loads the workbook and unmerges all merged cell ranges,
     copying the top-left value into every cell that was part of the merge.
  2. The cleaned workbook is written to an in-memory BytesIO buffer
     (zero disk I/O).
  3. pandas reads all sheets from the buffer.
  4. Each row is stringified as:
     "[Sheet: <name>] [<col1>: <val1>] [<col2>: <val2>] ..."
  5. Empty rows are dropped; NaN values are omitted from the chunk.

Why not just use pandas directly?
  Merged cells → pandas sees NaN in all but the top-left cell.
  Row-level stringification → blind splitters destroy column-header context.
"""

from __future__ import annotations

import io
import logging
from pathlib import Path
from typing import Dict, List, Optional

import pandas as pd
from openpyxl import load_workbook

from ..vision_client import GeminiVisionClient, DECORATIVE_MARKER

logger = logging.getLogger(__name__)


class ExcelParser:
    """
    Semantic chunker for Excel (.xlsx) workbooks.

    Each row in each sheet becomes an independent chunk with full
    column-header context and sheet name injected into raw_context.
    Images anchored to cells are triaged via GeminiVisionClient.

    Usage:
        parser = ExcelParser()
        chunks = parser.parse_file("/path/to/data.xlsx")
    """

    def __init__(
        self,
        vision_client: Optional[GeminiVisionClient] = None,
    ) -> None:
        self._vision_client = vision_client or GeminiVisionClient()

    # ── Public API ──────────────────────────────────────────────────

    def parse_file(self, file_path: str) -> List[Dict[str, str]]:
        """
        Read an .xlsx file, unmerge cells, and produce row-level chunks.

        Returns:
            List of dicts with keys: "raw_context", "source_file"
        """
        path = Path(file_path)
        logger.info("ExcelParser: Loading %s", path.name)

        # Step 1: Load with openpyxl and unmerge all cells
        wb = load_workbook(str(path))
        self._unmerge_all_cells(wb)

        # Step 1b: Extract image anchors per sheet BEFORE closing
        image_map = self._extract_image_map(wb)

        # Step 1c: Strip images from the workbook so wb.save() won't
        # try to re-serialise them (their internal BytesIO is closed).
        for ws in wb.worksheets:
            ws._images = []

        # Step 2: Write cleaned workbook to in-memory buffer
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        wb.close()

        # Step 3: Read all sheets with pandas
        sheets = pd.read_excel(buffer, sheet_name=None, engine="openpyxl")

        # Step 4: Stringify each row
        chunks: List[Dict[str, str]] = []
        for sheet_name, df in sheets.items():
            sheet_images = image_map.get(sheet_name, {})
            sheet_chunks = self._stringify_sheet(
                df, sheet_name, path.name,
                sheet_images=sheet_images,
                vision_client=self._vision_client,
            )
            chunks.extend(sheet_chunks)

        logger.info(
            "ExcelParser produced %d chunks from %s (%d sheets).",
            len(chunks),
            path.name,
            len(sheets),
        )
        return chunks

    def parse_buffer(
        self,
        buffer: io.BytesIO,
        source_file: str = "unknown.xlsx",
    ) -> List[Dict[str, str]]:
        """
        Parse from an in-memory buffer (useful for testing without disk I/O).
        """
        # Load, unmerge, re-save to new buffer
        buffer.seek(0)
        wb = load_workbook(buffer)
        self._unmerge_all_cells(wb)

        # Extract image anchors before closing
        image_map = self._extract_image_map(wb)

        # Strip images so save won't hit closed handles
        for ws in wb.worksheets:
            ws._images = []

        clean_buffer = io.BytesIO()
        wb.save(clean_buffer)
        clean_buffer.seek(0)
        wb.close()

        sheets = pd.read_excel(clean_buffer, sheet_name=None, engine="openpyxl")

        chunks: List[Dict[str, str]] = []
        for sheet_name, df in sheets.items():
            sheet_images = image_map.get(sheet_name, {})
            sheet_chunks = self._stringify_sheet(
                df, sheet_name, source_file,
                sheet_images=sheet_images,
                vision_client=self._vision_client,
            )
            chunks.extend(sheet_chunks)

        return chunks

    # ── Internal: Merged Cell Resolution ────────────────────────────

    @staticmethod
    def _unmerge_all_cells(wb) -> None:
        """
        Iterate every sheet in the workbook. For each merged range:
          1. Extract the value from the top-left cell.
          2. Unmerge the range.
          3. Write the extracted value into every cell that was merged.

        This ensures pandas sees the correct value in every cell,
        not NaN in the non-top-left positions.
        """
        for ws in wb.worksheets:
            # We must collect ranges first because unmerging while
            # iterating mutates the collection.
            merged_ranges = list(ws.merged_cells.ranges)

            for merge_range in merged_ranges:
                # Top-left cell holds the value
                min_row = merge_range.min_row
                min_col = merge_range.min_col
                top_left_value = ws.cell(row=min_row, column=min_col).value

                # Unmerge first (required by openpyxl before writing to merged cells)
                ws.unmerge_cells(str(merge_range))

                # Fill every cell in the previously merged range
                for row in range(merge_range.min_row, merge_range.max_row + 1):
                    for col in range(merge_range.min_col, merge_range.max_col + 1):
                        ws.cell(row=row, column=col).value = top_left_value

            if merged_ranges:
                logger.info(
                    "Unmerged %d cell ranges in sheet '%s'.",
                    len(merged_ranges),
                    ws.title,
                )

    # ── Internal: Image Extraction ──────────────────────────────────

    @staticmethod
    def _extract_image_map(wb) -> Dict[str, Dict[int, List[tuple]]]:
        """
        Scan all sheets for images and return a map of:
          { sheet_name: { row_number: [(col_letter, image_bytes), ...] } }

        The row_number is 0-indexed (pandas convention) and accounts for
        the header row by subtracting 2 from the openpyxl 1-indexed anchor.
        """
        from openpyxl.utils import get_column_letter

        image_map: Dict[str, Dict[int, List[tuple]]] = {}

        for ws in wb.worksheets:
            sheet_images: Dict[int, List[tuple]] = {}

            for img in ws._images:
                anchor = getattr(img, "anchor", None)
                if anchor is None:
                    continue

                # TwoCellAnchor / OneCellAnchor — top-left cell
                _from = getattr(anchor, "_from", None) or getattr(anchor, "from_", None)
                if _from is None:
                    continue

                # openpyxl uses 0-indexed row/col on the anchor _from
                anchor_row = _from.row       # 0-indexed
                anchor_col = _from.col + 1   # 1-indexed for get_column_letter

                # Adjust for pandas: pandas row 0 == openpyxl row 2
                # (row 1 is the header). So pandas_row = anchor_row - 1.
                pandas_row = anchor_row - 1

                col_letter = get_column_letter(anchor_col)

                # Extract the image bytes from the image blob
                image_data = img._data()

                if pandas_row not in sheet_images:
                    sheet_images[pandas_row] = []
                sheet_images[pandas_row].append((col_letter, image_data))

            if sheet_images:
                image_map[ws.title] = sheet_images

        return image_map

    # ── Internal: Row Stringification ───────────────────────────────

    @staticmethod
    def _stringify_sheet(
        df: pd.DataFrame,
        sheet_name: str,
        source_file: str,
        sheet_images: Optional[Dict[int, List[tuple]]] = None,
        vision_client: Optional[GeminiVisionClient] = None,
    ) -> List[Dict[str, str]]:
        """
        Convert a dataframe into row-level chunk dicts.

        Format: "[Sheet: <name>] [<col>: <val>] [<col>: <val>] ..."

        Rules:
          • Completely empty rows are dropped.
          • NaN values are omitted from the chunk string.
          • Column headers are cleaned (stripped of whitespace).
          • If an image is anchored to this row, it is triaged via
            the vision client and injected as [Image at Col X: ...].
        """
        if sheet_images is None:
            sheet_images = {}

        # Drop fully empty rows
        df = df.dropna(how="all").reset_index(drop=True)

        # Clean column names
        columns = [str(c).strip() for c in df.columns]

        chunks: List[Dict[str, str]] = []

        for row_idx, row in df.iterrows():
            parts: List[str] = [f"[Sheet: {sheet_name}]"]

            for col_name, value in zip(columns, row.values):
                # Skip NaN / None
                if pd.isna(value):
                    continue
                # Convert value to string and strip
                val_str = str(value).strip()
                if val_str:
                    parts.append(f"[{col_name}: {val_str}]")

            # A row with only the sheet header has no data — skip it
            if len(parts) <= 1:
                continue

            raw_context = " ".join(parts)

            # ── Inject image transcriptions for this row ────────────
            if row_idx in sheet_images and vision_client is not None:
                for col_letter, image_data in sheet_images[row_idx]:
                    transcription = vision_client.transcribe_image(
                        image_data,
                        surrounding_context=raw_context,
                    )
                    if transcription != DECORATIVE_MARKER:
                        raw_context += f" [Image at Col {col_letter}: {transcription}]"

            chunks.append({
                "raw_context": raw_context,
                "source_file": source_file,
            })

        return chunks
