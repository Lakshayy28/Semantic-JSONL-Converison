"""
Phase 7 Tests — Contextual Chunking & FastAPI Integration
==========================================================
Validates:
  1. GeminiContextClient payload schema and mock behavior
  2. Global context injection into every chunk via SemanticRouter
  3. FastAPI /health, /convert, and /convert/batch endpoints
  4. Unchunked mode does NOT inject global context (bypass)
  5. Context fallback behavior
"""

from __future__ import annotations

import io
import json
import os
import tempfile
import zipfile
from pathlib import Path
from typing import List
from unittest.mock import patch

import pytest

from semantic_pipeline.context_client import (
    GeminiContextClient,
    CONTEXT_SYSTEM_PROMPT,
    FALLBACK_STRING,
    MOCK_CONTEXT,
)
from semantic_pipeline.models import EmitterConfig
from semantic_pipeline.router import SemanticRouter


# ── Helpers ────────────────────────────────────────────────────────

def _create_md_file(content: str = None) -> str:
    """Create a temp Markdown file and return its path."""
    if content is None:
        content = """# Credit Policy

## Eligibility Rules
Applicants must have a credit score above 650.
Income must exceed three times the monthly payment.

## Approval Workflow
1. Submit application through the portal.
2. Automated scoring engine evaluates the risk.
3. Manual review for edge cases above $500k.
"""
    path = tempfile.mktemp(suffix=".md")
    Path(path).write_text(content, encoding="utf-8")
    return path


def _create_xlsx_file() -> str:
    """Create a temp Excel file and return its path."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Rules"
    ws["A1"] = "Rule ID"
    ws["B1"] = "Condition"
    ws["C1"] = "Action"
    ws["A2"] = "R001"
    ws["B2"] = "Score > 700"
    ws["C2"] = "Auto Approve"
    ws["A3"] = "R002"
    ws["B3"] = "Score < 500"
    ws["C3"] = "Reject"

    path = tempfile.mktemp(suffix=".xlsx")
    wb.save(path)
    wb.close()
    return path


# ═══════════════════════════════════════════════════════════════════
# TEST GROUP 1 — GeminiContextClient
# ═══════════════════════════════════════════════════════════════════


class TestContextClient:
    """Verify the GeminiContextClient payload and mock behavior."""

    def test_payload_model(self):
        client = GeminiContextClient()
        payload = client.prepare_payload("Some document text.")
        assert payload["model"] == "gemini-2.5-flash"

    def test_payload_has_system_prompt(self):
        client = GeminiContextClient()
        payload = client.prepare_payload("Some document text.")
        assert payload["messages"][0]["role"] == "system"
        assert payload["messages"][0]["content"] == CONTEXT_SYSTEM_PROMPT

    def test_payload_has_user_message(self):
        client = GeminiContextClient()
        payload = client.prepare_payload("Document content here.")
        assert payload["messages"][1]["role"] == "user"
        assert payload["messages"][1]["content"] == "Document content here."

    def test_payload_has_top_p(self):
        client = GeminiContextClient()
        payload = client.prepare_payload("text")
        assert payload["top_p"] == 1

    @patch.dict(os.environ, {"GEMINI_API_KEY": "MOCK"})
    def test_mock_returns_deterministic_string(self):
        client = GeminiContextClient()
        result = client.generate_global_context("Some document text.")
        assert result == MOCK_CONTEXT
        assert "MOCK_GLOBAL_CONTEXT" in result

    @patch.dict(os.environ, {"GEMINI_API_KEY": "MOCK"})
    def test_empty_text_returns_fallback(self):
        client = GeminiContextClient()
        result = client.generate_global_context("")
        assert result == FALLBACK_STRING

    def test_explicit_mock_key(self):
        client = GeminiContextClient(api_key="MOCK")
        result = client.generate_global_context("text")
        assert result == MOCK_CONTEXT

    def test_default_timeout(self):
        client = GeminiContextClient()
        assert client.timeout == 60

    def test_default_max_retries(self):
        client = GeminiContextClient()
        assert client.max_retries == 3


# ═══════════════════════════════════════════════════════════════════
# TEST GROUP 2 — Context Injection in SemanticRouter
# ═══════════════════════════════════════════════════════════════════


class TestContextInjection:
    """Verify that global context is prepended to every chunk."""

    @patch.dict(os.environ, {"GEMINI_API_KEY": "MOCK"})
    def test_md_chunks_have_global_context_prefix(self):
        """Every chunk from a chunked MD file should start with the global context tag."""
        md_path = _create_md_file()
        try:
            with tempfile.TemporaryDirectory() as tmpdir:
                config = EmitterConfig(
                    output_dir=tmpdir,
                    base_filename="ctx_test",
                    usecase_id="test",
                )
                with SemanticRouter(config) as router:
                    records = router.ingest_file(md_path)

                assert len(records) > 0
                for r in records:
                    assert r.raw_context.startswith("[Global Document Context: MOCK_GLOBAL_CONTEXT")
        finally:
            Path(md_path).unlink(missing_ok=True)

    @patch.dict(os.environ, {"GEMINI_API_KEY": "MOCK"})
    def test_xlsx_chunks_have_global_context_prefix(self):
        """Every chunk from a chunked Excel file should start with the global context tag."""
        xlsx_path = _create_xlsx_file()
        try:
            with tempfile.TemporaryDirectory() as tmpdir:
                config = EmitterConfig(
                    output_dir=tmpdir,
                    base_filename="ctx_test",
                    usecase_id="test",
                )
                with SemanticRouter(config) as router:
                    records = router.ingest_file(xlsx_path)

                assert len(records) > 0
                for r in records:
                    assert r.raw_context.startswith("[Global Document Context: MOCK_GLOBAL_CONTEXT")
        finally:
            Path(xlsx_path).unlink(missing_ok=True)

    @patch.dict(os.environ, {"GEMINI_API_KEY": "MOCK"})
    def test_global_context_tag_format(self):
        """Verify the exact tag format: [Global Document Context: ...]\\n\\n<content>"""
        md_path = _create_md_file()
        try:
            with tempfile.TemporaryDirectory() as tmpdir:
                config = EmitterConfig(
                    output_dir=tmpdir,
                    base_filename="ctx_test",
                    usecase_id="test",
                )
                with SemanticRouter(config) as router:
                    records = router.ingest_file(md_path)

                assert len(records) > 0
                first = records[0].raw_context
                # Must have the tag, then double newline, then the original content
                assert "[Global Document Context:" in first
                assert "]\n\n" in first
                # After the tag, original content should still be present
                after_tag = first.split("]\n\n", 1)[1]
                assert len(after_tag) > 0
        finally:
            Path(md_path).unlink(missing_ok=True)

    @patch.dict(os.environ, {"GEMINI_API_KEY": "MOCK"})
    def test_unchunked_mode_no_global_context(self):
        """Unchunked mode should NOT inject global context — it's the full doc."""
        md_path = _create_md_file()
        try:
            with tempfile.TemporaryDirectory() as tmpdir:
                config = EmitterConfig(
                    output_dir=tmpdir,
                    base_filename="ctx_test",
                    usecase_id="test",
                )
                with SemanticRouter(config) as router:
                    records = router.ingest_file_unchunked(md_path)

                assert len(records) == 1
                assert not records[0].raw_context.startswith("[Global Document Context:")
        finally:
            Path(md_path).unlink(missing_ok=True)

    @patch.dict(os.environ, {"GEMINI_API_KEY": "MOCK"})
    def test_original_content_preserved(self):
        """Global context injection should NOT remove original chunk content."""
        md_path = _create_md_file()
        try:
            with tempfile.TemporaryDirectory() as tmpdir:
                config = EmitterConfig(
                    output_dir=tmpdir,
                    base_filename="ctx_test",
                    usecase_id="test",
                )
                with SemanticRouter(config) as router:
                    records = router.ingest_file(md_path)

                combined = " ".join(r.raw_context for r in records)
                assert "credit score" in combined.lower()
                assert "650" in combined
        finally:
            Path(md_path).unlink(missing_ok=True)


# ═══════════════════════════════════════════════════════════════════
# TEST GROUP 3 — Context Client Fallback
# ═══════════════════════════════════════════════════════════════════


class TestContextFallback:
    """Verify graceful degradation when the context API fails."""

    def test_retry_on_failure(self):
        """Failures should degrade gracefully and return FALLBACK_STRING."""
        client = GeminiContextClient(api_key="REAL_KEY")
        call_count = 0

        def _failing(payload):
            nonlocal call_count
            call_count += 1
            raise TimeoutError("API timeout")

        client._execute_request = _failing
        result = client.generate_global_context("Some text")

        assert result == FALLBACK_STRING
        assert call_count == 3  # default max_retries

    def test_fallback_does_not_crash_router(self):
        """If context generation fails, the router should still produce chunks — just without context."""
        md_path = _create_md_file()
        try:
            failing_client = GeminiContextClient(api_key="REAL_KEY")
            failing_client._execute_request = lambda p: (_ for _ in ()).throw(
                TimeoutError("down")
            )

            with tempfile.TemporaryDirectory() as tmpdir:
                config = EmitterConfig(
                    output_dir=tmpdir,
                    base_filename="fallback",
                    usecase_id="test",
                )
                with SemanticRouter(config, context_client=failing_client) as router:
                    records = router.ingest_file(md_path)

                assert len(records) > 0
                # Since context generation failed, no global context prefix
                for r in records:
                    assert not r.raw_context.startswith("[Global Document Context:")
        finally:
            Path(md_path).unlink(missing_ok=True)


# ═══════════════════════════════════════════════════════════════════
# TEST GROUP 4 — FastAPI Endpoint Tests
# ═══════════════════════════════════════════════════════════════════


class TestFastAPIEndpoints:
    """Verify FastAPI endpoints match the spec using TestClient."""

    @pytest.fixture(autouse=True)
    def setup(self):
        """Import and configure TestClient."""
        from fastapi.testclient import TestClient
        import api as api_module
        self.client = TestClient(api_module.app)

    def test_health_endpoint(self):
        """GET /health should return status, version, and supported extensions."""
        resp = self.client.get("/health")
        assert resp.status_code == 200
        data = resp.json()
        assert data["status"] == "healthy"
        assert data["pipeline_version"] == "1.0.0"
        assert ".xlsx" in data["supported_extensions"]
        assert ".docx" in data["supported_extensions"]
        assert ".md" in data["supported_extensions"]

    @patch.dict(os.environ, {"GEMINI_API_KEY": "MOCK"})
    def test_convert_returns_jsonl_file(self):
        """POST /convert should return a downloadable JSONL file."""
        md_path = _create_md_file()
        try:
            with open(md_path, "rb") as f:
                resp = self.client.post(
                    "/convert?usecase_id=test&identifier=ci",
                    files={"file": ("test.md", f, "text/markdown")},
                )
            assert resp.status_code == 200
            assert "application/x-ndjson" in resp.headers["content-type"]

            # Validate JSONL content
            lines = resp.text.strip().split("\n")
            assert len(lines) >= 1
            first = json.loads(lines[0])
            assert "raw_context" in first
            assert "chunk_id" in first
            assert "usecase_id" in first
        finally:
            Path(md_path).unlink(missing_ok=True)

    @patch.dict(os.environ, {"GEMINI_API_KEY": "MOCK"})
    def test_convert_unchunked(self):
        """POST /convert?chunked=false should return single-record JSONL."""
        md_path = _create_md_file()
        try:
            with open(md_path, "rb") as f:
                resp = self.client.post(
                    "/convert?usecase_id=test&chunked=false",
                    files={"file": ("test.md", f, "text/markdown")},
                )
            assert resp.status_code == 200
            lines = resp.text.strip().split("\n")
            assert len(lines) == 1
        finally:
            Path(md_path).unlink(missing_ok=True)

    @patch.dict(os.environ, {"GEMINI_API_KEY": "MOCK"})
    def test_convert_batch_returns_zip(self):
        """POST /convert/batch should return a ZIP with individual JSONL files."""
        md_path1 = _create_md_file("# Doc A\nContent of doc A.")
        md_path2 = _create_md_file("# Doc B\nContent of doc B.")
        try:
            with open(md_path1, "rb") as f1, open(md_path2, "rb") as f2:
                resp = self.client.post(
                    "/convert/batch?usecase_id=test",
                    files=[
                        ("files", ("doc_a.md", f1, "text/markdown")),
                        ("files", ("doc_b.md", f2, "text/markdown")),
                    ],
                )
            assert resp.status_code == 200
            assert "application/zip" in resp.headers["content-type"]

            # Validate ZIP contents
            z = zipfile.ZipFile(io.BytesIO(resp.content))
            names = z.namelist()
            assert len(names) == 2
            assert "doc_a.jsonl" in names
            assert "doc_b.jsonl" in names

            # Validate each JSONL inside the ZIP
            for name in names:
                content = z.read(name).decode("utf-8")
                lines = content.strip().split("\n")
                assert len(lines) >= 1
                record = json.loads(lines[0])
                assert "raw_context" in record
        finally:
            Path(md_path1).unlink(missing_ok=True)
            Path(md_path2).unlink(missing_ok=True)

    def test_convert_unsupported_type(self):
        """POST /convert with unsupported file type should return 400."""
        content = b"not a real file"
        resp = self.client.post(
            "/convert",
            files={"file": ("test.exe", io.BytesIO(content), "application/octet-stream")},
        )
        assert resp.status_code == 400

    @patch.dict(os.environ, {"GEMINI_API_KEY": "MOCK"})
    def test_convert_chunked_has_global_context(self):
        """POST /convert (chunked=true) should produce chunks with global context."""
        md_path = _create_md_file()
        try:
            with open(md_path, "rb") as f:
                resp = self.client.post(
                    "/convert?usecase_id=test&chunked=true",
                    files={"file": ("test.md", f, "text/markdown")},
                )
            assert resp.status_code == 200
            lines = resp.text.strip().split("\n")
            first = json.loads(lines[0])
            assert first["raw_context"].startswith("[Global Document Context: MOCK_GLOBAL_CONTEXT")
        finally:
            Path(md_path).unlink(missing_ok=True)
