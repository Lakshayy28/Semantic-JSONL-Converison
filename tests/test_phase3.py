"""
Phase 3 Tests — Excel Tabular Parser
======================================
These tests verify:
  1. ExcelParser: merged cell unrolling, multi-sheet processing,
     row-level stringification with column-header context.
  2. Edge cases: empty rows, NaN handling, sheet name injection.
  3. SemanticRouter integration: .xlsx files produce ChunkRecords.

All workbooks are created programmatically in-memory using openpyxl
to avoid test dependency on external fixtures.
"""

from __future__ import annotations

import io
import json
from pathlib import Path
from typing import List

import pytest
from openpyxl import Workbook

from semantic_pipeline.models import ChunkRecord, EmitterConfig
from semantic_pipeline.parsers.excel_parser import ExcelParser
from semantic_pipeline.router import SemanticRouter


# ═══════════════════════════════════════════════════════════════════
# Fixtures — Workbook Generators
# ═══════════════════════════════════════════════════════════════════


@pytest.fixture
def excel_parser() -> ExcelParser:
    """Default ExcelParser instance."""
    return ExcelParser()


@pytest.fixture
def simple_workbook(tmp_path: Path) -> Path:
    """
    Create a simple Excel workbook with 2 sheets, each having
    a header row and 3 data rows.
    """
    wb = Workbook()

    # Sheet 1: Employees
    ws1 = wb.active
    ws1.title = "Employees"
    ws1.append(["Name", "Department", "Salary"])
    ws1.append(["Alice", "Engineering", 120000])
    ws1.append(["Bob", "Marketing", 95000])
    ws1.append(["Charlie", "Engineering", 110000])

    # Sheet 2: Projects
    ws2 = wb.create_sheet("Projects")
    ws2.append(["Project ID", "Name", "Status"])
    ws2.append(["P-001", "Data Pipeline", "Active"])
    ws2.append(["P-002", "Auth Service", "Complete"])

    path = tmp_path / "simple.xlsx"
    wb.save(str(path))
    wb.close()
    return path


@pytest.fixture
def merged_cells_workbook(tmp_path: Path) -> Path:
    """
    CRITICAL FIXTURE: Create a workbook where column A rows 2–4
    are merged with the value "Q3 Data". This simulates the
    enterprise pattern where a category label spans multiple rows.

    Layout after merge:
        A           B          C
    1   Quarter     Metric     Value
    2   Q3 Data     Revenue    500
    3   (merged)    Costs      200
    4   (merged)    Profit     300
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Quarterly"

    # Header row
    ws["A1"] = "Quarter"
    ws["B1"] = "Metric"
    ws["C1"] = "Value"

    # Data rows
    ws["A2"] = "Q3 Data"
    ws["B2"] = "Revenue"
    ws["C2"] = 500

    ws["A3"] = None  # Will be merged
    ws["B3"] = "Costs"
    ws["C3"] = 200

    ws["A4"] = None  # Will be merged
    ws["B4"] = "Profit"
    ws["C4"] = 300

    # Merge A2:A4 — this is the enterprise use case
    ws.merge_cells("A2:A4")

    path = tmp_path / "merged.xlsx"
    wb.save(str(path))
    wb.close()
    return path


@pytest.fixture
def multi_merge_workbook(tmp_path: Path) -> Path:
    """
    Workbook with multiple merged ranges across rows AND columns.

    Layout:
        A           B           C          D
    1   Region      Region      Product    Sales
    2   North       North       Widget     100
    3   North       North       Gadget     200
    4   South       South       Widget     150

    Merges: A1:B1 ("Region"), A2:B3 ("North"), A4:B4 ("South")
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "SalesData"

    ws["A1"] = "Region"
    ws["B1"] = "Region"
    ws["C1"] = "Product"
    ws["D1"] = "Sales"

    ws["A2"] = "North"
    ws["B2"] = None
    ws["C2"] = "Widget"
    ws["D2"] = 100

    ws["A3"] = None
    ws["B3"] = None
    ws["C3"] = "Gadget"
    ws["D3"] = 200

    ws["A4"] = "South"
    ws["B4"] = None
    ws["C4"] = "Widget"
    ws["D4"] = 150

    # Column merges
    ws.merge_cells("A1:B1")
    # Block merge: A2:B3
    ws.merge_cells("A2:B3")
    # Row merge: A4:B4
    ws.merge_cells("A4:B4")

    path = tmp_path / "multi_merge.xlsx"
    wb.save(str(path))
    wb.close()
    return path


@pytest.fixture
def empty_rows_workbook(tmp_path: Path) -> Path:
    """Workbook with some completely empty rows interspersed."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sparse"

    ws.append(["ID", "Name", "Score"])
    ws.append([1, "Alice", 95])
    ws.append([None, None, None])  # empty row
    ws.append([None, None, None])  # empty row
    ws.append([2, "Bob", 88])
    ws.append([None, None, None])  # empty row
    ws.append([3, "Charlie", 92])

    path = tmp_path / "sparse.xlsx"
    wb.save(str(path))
    wb.close()
    return path


@pytest.fixture
def nan_values_workbook(tmp_path: Path) -> Path:
    """Workbook with some cells containing NaN/None values."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Partial"

    ws.append(["Rule ID", "Action", "Notes"])
    ws.append([101, "Approve", None])  # Notes is NaN
    ws.append([102, None, "Needs review"])  # Action is NaN
    ws.append([103, "Decline", "High risk"])

    path = tmp_path / "partial.xlsx"
    wb.save(str(path))
    wb.close()
    return path


# ═══════════════════════════════════════════════════════════════════
# 1. Basic Parsing Tests
# ═══════════════════════════════════════════════════════════════════


class TestExcelParserBasic:
    """Validate fundamental row-level parsing and sheet awareness."""

    def test_simple_workbook_produces_chunks(
        self, excel_parser: ExcelParser, simple_workbook: Path
    ):
        """Each data row should produce one chunk."""
        chunks = excel_parser.parse_file(str(simple_workbook))
        # 3 rows in Employees + 2 rows in Projects = 5
        assert len(chunks) == 5

    def test_sheet_name_injected_into_raw_context(
        self, excel_parser: ExcelParser, simple_workbook: Path
    ):
        """Every chunk must contain [Sheet: <name>]."""
        chunks = excel_parser.parse_file(str(simple_workbook))

        employee_chunks = [c for c in chunks if "[Sheet: Employees]" in c["raw_context"]]
        project_chunks = [c for c in chunks if "[Sheet: Projects]" in c["raw_context"]]

        assert len(employee_chunks) == 3
        assert len(project_chunks) == 2

    def test_column_headers_mapped_to_values(
        self, excel_parser: ExcelParser, simple_workbook: Path
    ):
        """
        Each chunk should contain [ColumnHeader: CellValue] pairs.
        """
        chunks = excel_parser.parse_file(str(simple_workbook))

        # Find Alice's row
        alice_chunks = [c for c in chunks if "Alice" in c["raw_context"]]
        assert len(alice_chunks) == 1

        ctx = alice_chunks[0]["raw_context"]
        assert "[Name: Alice]" in ctx
        assert "[Department: Engineering]" in ctx
        assert "[Salary:" in ctx

    def test_source_file_propagated(
        self, excel_parser: ExcelParser, simple_workbook: Path
    ):
        """Every chunk should carry the source filename."""
        chunks = excel_parser.parse_file(str(simple_workbook))
        for c in chunks:
            assert c["source_file"] == "simple.xlsx"

    def test_raw_context_format(
        self, excel_parser: ExcelParser, simple_workbook: Path
    ):
        """
        Verify the exact format:
        "[Sheet: X] [Col1: Val1] [Col2: Val2] ..."
        """
        chunks = excel_parser.parse_file(str(simple_workbook))
        ctx = chunks[0]["raw_context"]

        assert ctx.startswith("[Sheet:")
        # Should have multiple bracketed key-value pairs
        bracket_count = ctx.count("[")
        assert bracket_count >= 3  # Sheet + at least 2 columns


# ═══════════════════════════════════════════════════════════════════
# 2. Merged Cell Tests (CRITICAL)
# ═══════════════════════════════════════════════════════════════════


class TestMergedCellResolution:
    """
    CRITICAL: Verify that merged cells are properly unrolled so that
    every row in the merge range receives the top-left cell's value.
    """

    def test_merged_value_propagated_to_all_rows(
        self, excel_parser: ExcelParser, merged_cells_workbook: Path
    ):
        """
        A2:A4 merged with "Q3 Data" — all 3 data rows must contain
        "Q3 Data" in their raw_context.
        """
        chunks = excel_parser.parse_file(str(merged_cells_workbook))
        assert len(chunks) == 3  # 3 data rows

        for chunk in chunks:
            assert "Q3 Data" in chunk["raw_context"], (
                f"Merged value 'Q3 Data' missing from chunk: "
                f"{chunk['raw_context'][:100]}"
            )

    def test_merged_value_paired_with_correct_header(
        self, excel_parser: ExcelParser, merged_cells_workbook: Path
    ):
        """The merged column value should be paired with its column header."""
        chunks = excel_parser.parse_file(str(merged_cells_workbook))

        for chunk in chunks:
            assert "[Quarter: Q3 Data]" in chunk["raw_context"]

    def test_non_merged_columns_independent(
        self, excel_parser: ExcelParser, merged_cells_workbook: Path
    ):
        """Columns B and C should have their own distinct values per row."""
        chunks = excel_parser.parse_file(str(merged_cells_workbook))

        metrics = [c["raw_context"] for c in chunks]
        assert any("[Metric: Revenue]" in m for m in metrics)
        assert any("[Metric: Costs]" in m for m in metrics)
        assert any("[Metric: Profit]" in m for m in metrics)

    def test_multi_dimensional_merge(
        self, excel_parser: ExcelParser, multi_merge_workbook: Path
    ):
        """
        Block merge (A2:B3 = "North") should fill all 4 cells.
        Row merge (A4:B4 = "South") should fill both columns.
        """
        chunks = excel_parser.parse_file(str(multi_merge_workbook))
        assert len(chunks) == 3  # 3 data rows

        # Rows 2 and 3 should both have "North"
        north_chunks = [c for c in chunks if "North" in c["raw_context"]]
        assert len(north_chunks) == 2

        # Row 4 should have "South"
        south_chunks = [c for c in chunks if "South" in c["raw_context"]]
        assert len(south_chunks) == 1


# ═══════════════════════════════════════════════════════════════════
# 3. Data Cleaning Tests
# ═══════════════════════════════════════════════════════════════════


class TestDataCleaning:
    """Verify empty row removal and NaN handling."""

    def test_empty_rows_dropped(
        self, excel_parser: ExcelParser, empty_rows_workbook: Path
    ):
        """Completely empty rows should not produce chunks."""
        chunks = excel_parser.parse_file(str(empty_rows_workbook))
        # Only 3 data rows (Alice, Bob, Charlie), not 6
        assert len(chunks) == 3

        names = " ".join(c["raw_context"] for c in chunks)
        assert "Alice" in names
        assert "Bob" in names
        assert "Charlie" in names

    def test_nan_values_omitted_from_context(
        self, excel_parser: ExcelParser, nan_values_workbook: Path
    ):
        """NaN cells should be omitted from the chunk string, not shown as 'nan'."""
        chunks = excel_parser.parse_file(str(nan_values_workbook))
        assert len(chunks) == 3

        for chunk in chunks:
            ctx = chunk["raw_context"]
            assert "nan" not in ctx.lower(), (
                f"NaN leaked into raw_context: {ctx}"
            )
            assert "None" not in ctx, (
                f"None leaked into raw_context: {ctx}"
            )

    def test_nan_column_omitted_not_placeholder(
        self, excel_parser: ExcelParser, nan_values_workbook: Path
    ):
        """
        Row 101 has Notes=None → the chunk should NOT contain [Notes: ...]
        Row 102 has Action=None → the chunk should NOT contain [Action: ...]
        """
        chunks = excel_parser.parse_file(str(nan_values_workbook))

        rule_101 = [c for c in chunks if "101" in c["raw_context"]][0]
        assert "[Notes:" not in rule_101["raw_context"]

        rule_102 = [c for c in chunks if "102" in c["raw_context"]][0]
        assert "[Action:" not in rule_102["raw_context"]


# ═══════════════════════════════════════════════════════════════════
# 4. In-Memory Buffer Parsing
# ═══════════════════════════════════════════════════════════════════


class TestBufferParsing:
    """Verify parse_buffer works without disk I/O."""

    def test_parse_from_buffer(self, excel_parser: ExcelParser):
        """Parsing from an in-memory BytesIO buffer should work identically."""
        wb = Workbook()
        ws = wb.active
        ws.title = "InMemory"
        ws.append(["Key", "Value"])
        ws.append(["host", "localhost"])
        ws.append(["port", "8080"])

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        wb.close()

        chunks = excel_parser.parse_buffer(buf, source_file="config.xlsx")
        assert len(chunks) == 2
        assert "[Sheet: InMemory]" in chunks[0]["raw_context"]
        assert chunks[0]["source_file"] == "config.xlsx"


# ═══════════════════════════════════════════════════════════════════
# 5. Router Integration Tests
# ═══════════════════════════════════════════════════════════════════


class TestSemanticRouterPhase3:
    """Verify the router now processes .xlsx files end-to-end."""

    def test_router_processes_xlsx(
        self, tmp_path: Path, simple_workbook: Path
    ):
        """Router should produce ChunkRecords from an .xlsx file."""
        output_dir = tmp_path / "output"
        config = EmitterConfig(
            output_dir=str(output_dir),
            usecase_id="credit-rules",
            identifier="underwriting",
        )

        with SemanticRouter(config) as router:
            chunks = router.ingest_file(str(simple_workbook))

        assert len(chunks) == 5
        assert all(isinstance(c, ChunkRecord) for c in chunks)
        assert all(c.usecase_id == "credit-rules" for c in chunks)
        assert all(c.file_name == "simple.xlsx" for c in chunks)

    def test_router_xlsx_writes_valid_jsonl(
        self, tmp_path: Path, simple_workbook: Path
    ):
        """Router output JSONL should contain valid JSON with schema fields."""
        output_dir = tmp_path / "output"
        config = EmitterConfig(output_dir=str(output_dir))

        with SemanticRouter(config) as router:
            router.ingest_file(str(simple_workbook))

        jsonl_files = list(output_dir.glob("*.jsonl"))
        assert len(jsonl_files) >= 1

        for f in jsonl_files:
            for line in f.read_text().strip().splitlines():
                obj = json.loads(line)
                assert "raw_context" in obj
                assert "[Sheet:" in obj["raw_context"]
                assert "chunk_id" in obj

    def test_router_xlsx_with_merges(
        self, tmp_path: Path, merged_cells_workbook: Path
    ):
        """Router should correctly handle merged cells in .xlsx files."""
        output_dir = tmp_path / "output"
        config = EmitterConfig(output_dir=str(output_dir))

        with SemanticRouter(config) as router:
            chunks = router.ingest_file(str(merged_cells_workbook))

        assert len(chunks) == 3
        # All chunks should have the merged value
        for c in chunks:
            assert "Q3 Data" in c.raw_context

    def test_router_summary_includes_xlsx(
        self, tmp_path: Path, simple_workbook: Path
    ):
        """Summary should list .xlsx as processed."""
        output_dir = tmp_path / "output"
        config = EmitterConfig(output_dir=str(output_dir))

        with SemanticRouter(config) as router:
            router.ingest_file(str(simple_workbook))
            summary = router.summary()

        assert len(summary["processed_files"]) == 1
        assert "simple.xlsx" in summary["processed_files"][0]

    def test_openapi_still_stubbed(self, tmp_path: Path):
        """.json/.yaml files should still be skipped until Phase 4."""
        json_file = tmp_path / "api.json"
        json_file.write_text('{"openapi": "3.0.0"}')

        output_dir = tmp_path / "output"
        config = EmitterConfig(output_dir=str(output_dir))

        with SemanticRouter(config) as router:
            chunks = router.ingest_file(str(json_file))

        assert chunks == []


# ═══════════════════════════════════════════════════════════════════
# 6. Real Artifact Test
# ═══════════════════════════════════════════════════════════════════


class TestRealArtifactExcel:
    """
    Test against the actual decision_controller_rules.xlsx artifact
    in the workspace if it exists.
    """

    ARTIFACT_PATH = Path(
        "/Users/lakshaychandra/Documents/Semantic JSONL Converison"
        "/artifacts/decision_controller_rules.xlsx"
    )

    @pytest.mark.skipif(
        not ARTIFACT_PATH.exists(),
        reason="Real artifact not available in workspace.",
    )
    def test_real_excel_artifact_parses(self, excel_parser: ExcelParser):
        """The actual enterprise Excel artifact should parse without errors."""
        chunks = excel_parser.parse_file(str(self.ARTIFACT_PATH))
        assert len(chunks) > 0, "Expected chunks from the real Excel artifact."

        # Every chunk should have the sheet and column context
        for c in chunks:
            assert "[Sheet:" in c["raw_context"]
            assert c["source_file"] == "decision_controller_rules.xlsx"
