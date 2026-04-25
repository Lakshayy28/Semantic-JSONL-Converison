"""
Phase 5 Tests — Context Hardening & Unchunked Bypass
======================================================
These tests verify:
  1. Unchunked bypass mode: the router aggregates all parser output into
     exactly one ChunkRecord per file, with chunk_id = '{doc_id}-full'.
  2. OpenAPI $ref resolution: jsonref inlines $ref pointers before
     endpoint traversal, so schema fields appear in raw_context.
  3. Word document table extraction: tables in DOCX files are converted
     to Markdown table syntax (| Header | ... | with separator row).
  4. Full backward compatibility: Phases 1–4 remain green.
"""

from __future__ import annotations
import os

import json
from pathlib import Path

import pytest
from docx import Document as DocxDocument
from openpyxl import Workbook

from semantic_pipeline.models import EmitterConfig
from semantic_pipeline.parsers.structured_parser import StructuredParser
from semantic_pipeline.parsers.word_parser import WordParser
from semantic_pipeline.router import SemanticRouter


# ═══════════════════════════════════════════════════════════════════
# Fixtures
# ═══════════════════════════════════════════════════════════════════


@pytest.fixture
def multi_header_md(tmp_path: Path) -> Path:
    """Markdown file with multiple headers and body sections."""
    content = """\
# Chapter 1

This is the introduction with some details about the system.

## 1.1 Overview

The system handles credit decisions in real time.

## 1.2 Components

There are three main services: Gateway, Engine, and Reporter.

# Chapter 2

This chapter covers deployment and operations.

## 2.1 Infrastructure

We use Kubernetes with auto-scaling policies.

## 2.2 Monitoring

Prometheus and Grafana provide observability.
"""
    f = tmp_path / "multi_header.md"
    f.write_text(content)
    return f


@pytest.fixture
def multi_sheet_xlsx(tmp_path: Path) -> Path:
    """Excel workbook with 2 sheets and multiple rows."""
    wb = Workbook()

    ws1 = wb.active
    ws1.title = "Rules"
    ws1.append(["Rule ID", "Threshold", "Action"])
    ws1.append([1, 700, "Approve"])
    ws1.append([2, 600, "Review"])
    ws1.append([3, 500, "Decline"])

    ws2 = wb.create_sheet("Overrides")
    ws2.append(["Override ID", "Reason", "Approver"])
    ws2.append(["OV-001", "VIP Customer", "Manager"])
    ws2.append(["OV-002", "Policy Exception", "Director"])

    path = tmp_path / "multi_sheet.xlsx"
    wb.save(str(path))
    wb.close()
    return path


@pytest.fixture
def openapi_spec_with_refs() -> dict:
    """
    OpenAPI 3.0 spec with $ref pointers to components/schemas.
    This is the CRITICAL test for $ref resolution.
    """
    return {
        "openapi": "3.0.0",
        "info": {"title": "Applicant API", "version": "1.0"},
        "components": {
            "schemas": {
                "Applicant": {
                    "type": "object",
                    "required": ["full_name", "ssn"],
                    "properties": {
                        "full_name": {"type": "string"},
                        "ssn": {"type": "string"},
                        "income": {"type": "number"},
                        "credit_score": {"type": "integer"},
                    },
                },
                "Decision": {
                    "type": "object",
                    "properties": {
                        "status": {"type": "string"},
                        "reason": {"type": "string"},
                    },
                },
            }
        },
        "paths": {
            "/applicants": {
                "post": {
                    "summary": "Submit a new applicant",
                    "tags": ["Applicants"],
                    "requestBody": {
                        "required": True,
                        "description": "Applicant data",
                        "content": {
                            "application/json": {
                                "schema": {
                                    "$ref": "#/components/schemas/Applicant"
                                }
                            }
                        },
                    },
                    "responses": {
                        "201": {
                            "description": "Created",
                            "content": {
                                "application/json": {
                                    "schema": {
                                        "$ref": "#/components/schemas/Decision"
                                    }
                                }
                            },
                        }
                    },
                }
            },
            "/applicants/{id}": {
                "get": {
                    "summary": "Get applicant details",
                    "tags": ["Applicants"],
                    "parameters": [
                        {
                            "name": "id",
                            "in": "path",
                            "required": True,
                            "schema": {"type": "string"},
                        }
                    ],
                    "responses": {
                        "200": {
                            "description": "Applicant found",
                            "content": {
                                "application/json": {
                                    "schema": {
                                        "$ref": "#/components/schemas/Applicant"
                                    }
                                }
                            },
                        }
                    },
                }
            },
        },
    }


@pytest.fixture
def openapi_ref_file(tmp_path: Path, openapi_spec_with_refs: dict) -> Path:
    """Write the ref-containing spec to a JSON file."""
    f = tmp_path / "ref_spec.json"
    f.write_text(json.dumps(openapi_spec_with_refs, indent=2))
    return f


@pytest.fixture
def docx_with_table(tmp_path: Path) -> Path:
    """Create a DOCX file containing a heading, paragraph, and a table."""
    doc = DocxDocument()
    doc.add_heading("Risk Assessment", level=1)
    doc.add_paragraph("The following table summarizes the risk categories:")

    table = doc.add_table(rows=4, cols=3)
    table.style = "Table Grid"

    # Header row
    table.cell(0, 0).text = "Category"
    table.cell(0, 1).text = "Threshold"
    table.cell(0, 2).text = "Action"

    # Data rows
    table.cell(1, 0).text = "Low"
    table.cell(1, 1).text = "< 300"
    table.cell(1, 2).text = "Auto Approve"

    table.cell(2, 0).text = "Medium"
    table.cell(2, 1).text = "300-700"
    table.cell(2, 2).text = "Manual Review"

    table.cell(3, 0).text = "High"
    table.cell(3, 1).text = "> 700"
    table.cell(3, 2).text = "Auto Decline"

    doc.add_paragraph("Additional rules may apply based on region.")

    path = tmp_path / "risk_assessment.docx"
    doc.save(str(path))
    return path


# ═══════════════════════════════════════════════════════════════════
# 1. Unchunked Bypass Mode — Markdown
# ═══════════════════════════════════════════════════════════════════


class TestUnchunkedMarkdown:
    """Verify whole-file mode with Markdown input."""

    def test_exactly_one_chunk_returned(
        self, tmp_path: Path, multi_header_md: Path
    ):
        """Unchunked mode must return exactly 1 ChunkRecord."""
        output_dir = tmp_path / "output"
        config = EmitterConfig(output_dir=str(output_dir))

        with SemanticRouter(config) as router:
            chunks = router.ingest_file_unchunked(str(multi_header_md))

        assert len(chunks) == 1

    def test_chunk_id_contains_full_suffix(
        self, tmp_path: Path, multi_header_md: Path
    ):
        """The chunk_id should end with '-full'."""
        output_dir = tmp_path / "output"
        config = EmitterConfig(output_dir=str(output_dir))

        with SemanticRouter(config) as router:
            chunks = router.ingest_file_unchunked(str(multi_header_md))

        assert chunks[0].chunk_id.endswith("-full")
        assert chunks[0].document_id in chunks[0].chunk_id

    def test_raw_context_contains_all_content(
        self, tmp_path: Path, multi_header_md: Path
    ):
        """The single raw_context must contain text from ALL sections."""
        output_dir = tmp_path / "output"
        config = EmitterConfig(output_dir=str(output_dir))

        with SemanticRouter(config) as router:
            chunks = router.ingest_file_unchunked(str(multi_header_md))

        ctx = chunks[0].raw_context
        assert "Chapter 1" in ctx
        assert "Chapter 2" in ctx
        assert "Overview" in ctx
        assert "Components" in ctx
        assert "Infrastructure" in ctx
        assert "Monitoring" in ctx
        assert "Prometheus" in ctx

    def test_chunked_vs_unchunked_difference(
        self, tmp_path: Path, multi_header_md: Path
    ):
        """
        Chunked mode should produce multiple chunks;
        unchunked mode should produce exactly 1.
        """
        config_chunked = EmitterConfig(output_dir=str(tmp_path / "chunked"))
        config_unchunked = EmitterConfig(output_dir=str(tmp_path / "unchunked"))

        with SemanticRouter(config_chunked) as router:
            chunked = router.ingest_file(str(multi_header_md))

        with SemanticRouter(config_unchunked) as router:
            unchunked = router.ingest_file_unchunked(str(multi_header_md))

        assert len(chunked) > 1, "Chunked mode should produce multiple chunks"
        assert len(unchunked) == 1, "Unchunked mode must produce exactly 1"


# ═══════════════════════════════════════════════════════════════════
# 2. Unchunked Bypass Mode — Excel
# ═══════════════════════════════════════════════════════════════════


class TestUnchunkedExcel:
    """Verify whole-file mode with Excel input."""

    def test_exactly_one_chunk_from_multi_sheet(
        self, tmp_path: Path, multi_sheet_xlsx: Path
    ):
        """Multi-sheet workbook in unchunked mode → exactly 1 record."""
        output_dir = tmp_path / "output"
        config = EmitterConfig(output_dir=str(output_dir))

        with SemanticRouter(config) as router:
            chunks = router.ingest_file_unchunked(str(multi_sheet_xlsx))

        assert len(chunks) == 1

    def test_raw_context_contains_all_sheets(
        self, tmp_path: Path, multi_sheet_xlsx: Path
    ):
        """The single record must contain data from BOTH sheets."""
        output_dir = tmp_path / "output"
        config = EmitterConfig(output_dir=str(output_dir))

        with SemanticRouter(config) as router:
            chunks = router.ingest_file_unchunked(str(multi_sheet_xlsx))

        ctx = chunks[0].raw_context
        assert "[Sheet: Rules]" in ctx
        assert "[Sheet: Overrides]" in ctx
        assert "Approve" in ctx
        assert "VIP Customer" in ctx

    def test_unchunked_emits_to_jsonl(
        self, tmp_path: Path, multi_sheet_xlsx: Path
    ):
        """Unchunked records should still be written to JSONL output."""
        output_dir = tmp_path / "output"
        config = EmitterConfig(output_dir=str(output_dir))

        with SemanticRouter(config) as router:
            router.ingest_file_unchunked(str(multi_sheet_xlsx))

        jsonl_files = list(output_dir.glob("*.jsonl"))
        assert len(jsonl_files) >= 1

        lines = jsonl_files[0].read_text().strip().splitlines()
        assert len(lines) == 1  # exactly 1 record in the file

        obj = json.loads(lines[0])
        assert "raw_context" in obj
        assert obj["chunk_id"].endswith("-full")


# ═══════════════════════════════════════════════════════════════════
# 3. Unchunked Bypass Mode — Edge Cases
# ═══════════════════════════════════════════════════════════════════


class TestUnchunkedEdgeCases:
    """Edge cases for the unchunked bypass mode."""

    def test_unsupported_extension_returns_empty(self, tmp_path: Path):
        """Unsupported file types should return empty, not crash."""
        output_dir = tmp_path / "output"
        config = EmitterConfig(output_dir=str(output_dir))

        (tmp_path / "test.csv").write_text("a,b,c\n1,2,3\n")

        with SemanticRouter(config) as router:
            chunks = router.ingest_file_unchunked(str(tmp_path / "test.csv"))

        assert chunks == []

    def test_unchunked_json(self, tmp_path: Path):
        """Generic JSON in unchunked mode should produce 1 record."""
        output_dir = tmp_path / "output"
        config = EmitterConfig(output_dir=str(output_dir))

        (tmp_path / "config.json").write_text('{"key": "value", "num": 42}')

        with SemanticRouter(config) as router:
            chunks = router.ingest_file_unchunked(str(tmp_path / "config.json"))

        assert len(chunks) == 1
        assert chunks[0].chunk_id.endswith("-full")

    def test_unchunked_summary_reflects_processed(
        self, tmp_path: Path, multi_header_md: Path
    ):
        """Summary should list unchunked files as processed."""
        output_dir = tmp_path / "output"
        config = EmitterConfig(output_dir=str(output_dir))

        with SemanticRouter(config) as router:
            router.ingest_file_unchunked(str(multi_header_md))
            summary = router.summary()

        assert len(summary["processed_files"]) == 1


# ═══════════════════════════════════════════════════════════════════
# 4. OpenAPI $ref Resolution
# ═══════════════════════════════════════════════════════════════════


class TestRefResolution:
    """
    CRITICAL: Verify that $ref pointers are resolved before endpoint
    traversal so that schema fields appear in raw_context.
    """

    def test_ref_resolved_schema_fields_appear(
        self, openapi_ref_file: Path
    ):
        """
        The Applicant schema has fields: full_name, ssn, income, credit_score.
        After $ref resolution, these must appear in the POST /applicants chunk.
        """
        parser = StructuredParser()
        chunks = parser.parse_file(str(openapi_ref_file))

        post_chunk = [c for c in chunks if "POST /applicants." in c["raw_context"]][0]
        ctx = post_chunk["raw_context"]

        # Schema fields must be inlined, NOT the literal "$ref" string
        assert "full_name" in ctx
        assert "ssn" in ctx
        assert "income" in ctx
        assert "credit_score" in ctx
        assert "Schema Fields:" in ctx

    def test_ref_string_not_in_output(self, openapi_ref_file: Path):
        """The literal string '$ref' should NOT appear in any chunk."""
        parser = StructuredParser()
        chunks = parser.parse_file(str(openapi_ref_file))

        for c in chunks:
            assert "$ref" not in c["raw_context"], (
                f"Unresolved $ref found in: {c['raw_context'][:200]}"
            )

    def test_ref_resolved_required_fields_tagged(
        self, openapi_ref_file: Path
    ):
        """
        The Applicant schema marks full_name and ssn as required.
        These should appear with 'required' tag in the context.
        """
        parser = StructuredParser()
        chunks = parser.parse_file(str(openapi_ref_file))

        post_chunk = [c for c in chunks if "POST /applicants." in c["raw_context"]][0]
        ctx = post_chunk["raw_context"]

        assert "full_name (string, required)" in ctx
        assert "ssn (string, required)" in ctx
        # income and credit_score are NOT required
        assert "income (number)" in ctx
        assert "credit_score (integer)" in ctx

    def test_ref_resolution_does_not_break_basic_parsing(
        self, openapi_ref_file: Path
    ):
        """Ref resolution should not interfere with basic endpoint parsing."""
        parser = StructuredParser()
        chunks = parser.parse_file(str(openapi_ref_file))

        # 2 endpoints: POST /applicants, GET /applicants/{id}
        assert len(chunks) == 2

        paths = [c["raw_context"] for c in chunks]
        assert any("POST /applicants." in p for p in paths)
        assert any("GET /applicants/{id}." in p for p in paths)

    def test_ref_resolution_via_parse_text(self, openapi_spec_with_refs: dict):
        """parse_text should also resolve $refs."""
        parser = StructuredParser()
        text = json.dumps(openapi_spec_with_refs)

        chunks = parser.parse_text(text, source_file="inline_ref.json", ext=".json")

        post_chunk = [c for c in chunks if "POST" in c["raw_context"]][0]
        assert "full_name" in post_chunk["raw_context"]

    def test_spec_without_refs_still_works(self, tmp_path: Path):
        """Specs without any $ref should parse normally (no crash)."""
        spec = {
            "openapi": "3.0.0",
            "info": {"title": "Simple", "version": "1.0"},
            "paths": {
                "/ping": {
                    "get": {
                        "summary": "Ping",
                        "responses": {"200": {"description": "pong"}},
                    }
                }
            },
        }
        f = tmp_path / "no_ref.json"
        f.write_text(json.dumps(spec))

        parser = StructuredParser()
        chunks = parser.parse_file(str(f))

        assert len(chunks) == 1
        assert "GET /ping" in chunks[0]["raw_context"]


# ═══════════════════════════════════════════════════════════════════
# 5. Word Document Table Extraction
# ═══════════════════════════════════════════════════════════════════


class TestWordTableExtraction:
    """Verify that Word tables are converted to Markdown table syntax."""

    def test_table_appears_in_raw_context(self, docx_with_table: Path):
        """Tables in DOCX should produce Markdown table syntax in chunks."""
        parser = WordParser()
        chunks = parser.parse_file(str(docx_with_table))

        all_context = " ".join(c["raw_context"] for c in chunks)
        # Markdown table markers
        assert "|" in all_context
        assert "---" in all_context

    def test_table_header_row_present(self, docx_with_table: Path):
        """The table header row should contain column names."""
        parser = WordParser()
        chunks = parser.parse_file(str(docx_with_table))

        all_context = " ".join(c["raw_context"] for c in chunks)
        assert "Category" in all_context
        assert "Threshold" in all_context
        assert "Action" in all_context

    def test_table_data_rows_present(self, docx_with_table: Path):
        """Data row values should appear in the raw_context."""
        parser = WordParser()
        chunks = parser.parse_file(str(docx_with_table))

        all_context = " ".join(c["raw_context"] for c in chunks)
        assert "Low" in all_context
        assert "Auto Approve" in all_context
        assert "Medium" in all_context
        assert "Manual Review" in all_context
        assert "High" in all_context
        assert "Auto Decline" in all_context

    def test_table_not_mashed_into_paragraph(self, docx_with_table: Path):
        """
        Table content should be structured (with | separators),
        NOT mashed into a flat paragraph string.
        """
        parser = WordParser()
        chunks = parser.parse_file(str(docx_with_table))

        all_context = "\n".join(c["raw_context"] for c in chunks)
        # A properly formatted Markdown table row looks like: | Val | Val | Val |
        assert "| Category |" in all_context or "| Low |" in all_context

    def test_non_table_content_preserved(self, docx_with_table: Path):
        """Paragraphs before and after the table should still appear."""
        parser = WordParser()
        chunks = parser.parse_file(str(docx_with_table))

        all_context = " ".join(c["raw_context"] for c in chunks)
        assert "Risk Assessment" in all_context
        assert "risk categories" in all_context
        assert "Additional rules" in all_context

    def test_table_to_markdown_static(self):
        """Test the _table_to_markdown static method directly."""
        doc = DocxDocument()
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Name"
        table.cell(0, 1).text = "Score"
        table.cell(1, 0).text = "Alice"
        table.cell(1, 1).text = "95"

        result = WordParser._table_to_markdown(table)
        assert "| Name | Score |" in result
        assert "| --- | --- |" in result
        assert "| Alice | 95 |" in result


# ═══════════════════════════════════════════════════════════════════
# 6. Router Integration — Unchunked + Ref + Tables Together
# ═══════════════════════════════════════════════════════════════════


class TestPhase5Integration:
    """End-to-end integration tests combining all Phase 5 features."""

    def test_unchunked_openapi_with_refs(
        self, tmp_path: Path, openapi_ref_file: Path
    ):
        """Unchunked mode with $ref spec should produce 1 record with all endpoints."""
        output_dir = tmp_path / "output"
        config = EmitterConfig(output_dir=str(output_dir))

        with SemanticRouter(config) as router:
            chunks = router.ingest_file_unchunked(str(openapi_ref_file))

        assert len(chunks) == 1
        ctx = chunks[0].raw_context
        # Both endpoints should be in the single record
        assert "POST /applicants" in ctx
        assert "GET /applicants/{id}" in ctx
        # Resolved schema fields should also be present
        assert "full_name" in ctx
        assert "credit_score" in ctx

    def test_unchunked_docx_with_table(
        self, tmp_path: Path, docx_with_table: Path
    ):
        """Unchunked DOCX with tables should have table markdown in one record."""
        output_dir = tmp_path / "output"
        config = EmitterConfig(output_dir=str(output_dir))

        with SemanticRouter(config) as router:
            chunks = router.ingest_file_unchunked(str(docx_with_table))

        assert len(chunks) == 1
        ctx = chunks[0].raw_context
        assert "Risk Assessment" in ctx
        assert "|" in ctx  # Markdown table syntax present
        assert "Auto Approve" in ctx

    def test_mixed_chunked_and_unchunked(
        self, tmp_path: Path, multi_header_md: Path, multi_sheet_xlsx: Path
    ):
        """
        You can mix chunked and unchunked calls within the same router session.
        """
        output_dir = tmp_path / "output"
        config = EmitterConfig(output_dir=str(output_dir))

        with SemanticRouter(config) as router:
            chunked = router.ingest_file(str(multi_header_md))
            unchunked = router.ingest_file_unchunked(str(multi_sheet_xlsx))
            summary = router.summary()

        assert len(chunked) > 1
        assert len(unchunked) == 1
        assert len(summary["processed_files"]) == 2
