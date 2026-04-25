"""
Phase 2 Tests — Markdown & Word Document Parsers
==================================================
These tests verify:
  1. MarkdownParser: header-aware chunking, breadcrumb stitching,
     recursive fallback for large sections, plain text handling.
  2. WordParser: DOCX heading style → Markdown conversion,
     full round-trip through MarkdownParser.
  3. SemanticRouter integration: .md and .docx files now produce
     real ChunkRecords instead of empty stubs.
"""

from __future__ import annotations
import os

import json
from pathlib import Path

import pytest
from docx import Document

from semantic_pipeline.models import ChunkRecord, EmitterConfig
from semantic_pipeline.parsers.markdown_parser import MarkdownParser
from semantic_pipeline.parsers.word_parser import WordParser
from semantic_pipeline.router import SemanticRouter


# ═══════════════════════════════════════════════════════════════════
# Fixtures
# ═══════════════════════════════════════════════════════════════════


@pytest.fixture
def md_parser() -> MarkdownParser:
    """Default MarkdownParser instance."""
    return MarkdownParser(chunk_size=2000, chunk_overlap=200)


@pytest.fixture
def word_parser() -> WordParser:
    """Default WordParser instance."""
    return WordParser(chunk_size=2000, chunk_overlap=200)


@pytest.fixture
def simple_markdown() -> str:
    """A well-structured Markdown document with nested headers."""
    return """\
# Architecture Overview

This document describes the system architecture.

## Components

### API Gateway

The gateway handles authentication, rate limiting, and request routing.

### Database Layer

PostgreSQL is used for persistent storage with connection pooling.

## Deployment

The system is deployed on Kubernetes with multi-AZ redundancy.
"""


@pytest.fixture
def massive_section_markdown() -> str:
    """
    Markdown with one section that exceeds the default 2000-char chunk_size,
    forcing the RecursiveCharacterTextSplitter fallback.
    """
    huge_body = "This is a critical compliance paragraph. " * 120  # ~5040 chars
    return f"""\
# Compliance Manual

## Section A: Overview

Short intro paragraph.

## Section B: Detailed Regulations

{huge_body}

## Section C: Summary

Brief closing remarks.
"""


@pytest.fixture
def sample_docx(tmp_path: Path) -> Path:
    """Generate a test .docx file with heading hierarchy using python-docx."""
    doc = Document()

    doc.add_heading("Business Architecture", level=1)
    doc.add_paragraph(
        "The Credit Decisioning Microservice acts as the central brain "
        "for automated underwriting workflows."
    )

    doc.add_heading("Core Use Cases", level=2)
    doc.add_paragraph(
        "Customers can request a credit line increase via the mobile app."
    )

    doc.add_heading("Credit Line Increase", level=3)
    doc.add_paragraph(
        "This triggers a real-time call to the Decisioning Engine "
        "which evaluates risk profiles."
    )

    doc.add_heading("Automated Risk Mitigation", level=3)
    doc.add_paragraph(
        "During economic downturns the microservice runs batch rules "
        "against high-risk portfolios."
    )

    doc.add_heading("Compliance", level=2)
    doc.add_paragraph(
        "Every decision yields a unique trace ID for complete explainability."
    )

    docx_path = tmp_path / "test_architecture.docx"
    doc.save(str(docx_path))
    return docx_path


@pytest.fixture
def docx_with_massive_section(tmp_path: Path) -> Path:
    """Generate a DOCX with one heading section that exceeds chunk_size."""
    doc = Document()
    doc.add_heading("Policy Document", level=1)
    doc.add_paragraph("Introduction to policies.")

    doc.add_heading("Detailed Regulations", level=2)
    huge_text = "All employees must comply with section 42 of the handbook. " * 100
    doc.add_paragraph(huge_text)

    doc.add_heading("Appendix", level=2)
    doc.add_paragraph("Reference tables and glossary.")

    path = tmp_path / "massive_section.docx"
    doc.save(str(path))
    return path


# ═══════════════════════════════════════════════════════════════════
# 1. MarkdownParser Tests
# ═══════════════════════════════════════════════════════════════════


class TestMarkdownParser:
    """Validate header-aware semantic chunking of Markdown content."""

    def test_basic_header_splitting(
        self, md_parser: MarkdownParser, simple_markdown: str
    ):
        """Each header section should produce a separate chunk."""
        chunks = md_parser.parse_text(simple_markdown, source_file="arch.md")
        assert len(chunks) >= 4  # At least: Overview, API Gateway, Database, Deployment

    def test_headers_stitched_into_raw_context(
        self, md_parser: MarkdownParser, simple_markdown: str
    ):
        """
        CRITICAL: Header metadata must appear inside raw_context,
        not just in metadata. The embedding API only sees raw_context.
        """
        chunks = md_parser.parse_text(simple_markdown, source_file="arch.md")

        # Find the chunk about API Gateway — it should contain the
        # full breadcrumb path: Architecture Overview > Components > API Gateway
        gateway_chunks = [
            c for c in chunks if "API Gateway" in c["raw_context"]
        ]
        assert len(gateway_chunks) >= 1, (
            "Expected at least one chunk mentioning 'API Gateway'."
        )

        ctx = gateway_chunks[0]["raw_context"]
        # Verify breadcrumb structure
        assert "Section:" in ctx, "Missing 'Section:' breadcrumb prefix."
        assert "Content:" in ctx, "Missing 'Content:' label."

    def test_nested_headers_produce_breadcrumb_path(
        self, md_parser: MarkdownParser, simple_markdown: str
    ):
        """
        A ### header under ## under # should produce a 3-level breadcrumb.
        Example: "Section: Architecture Overview > Components > API Gateway."
        """
        chunks = md_parser.parse_text(simple_markdown, source_file="arch.md")
        gateway_chunks = [
            c for c in chunks if "API Gateway" in c["raw_context"]
        ]
        ctx = gateway_chunks[0]["raw_context"]

        # Verify the full breadcrumb chain
        assert ">" in ctx, "Expected '>' delimiter in breadcrumb path."

    def test_recursive_fallback_for_massive_section(
        self, massive_section_markdown: str,
    ):
        """
        A section exceeding chunk_size must be sub-split by
        RecursiveCharacterTextSplitter, with headers preserved
        on EVERY sub-chunk.
        """
        parser = MarkdownParser(chunk_size=2000, chunk_overlap=200)
        chunks = parser.parse_text(
            massive_section_markdown, source_file="compliance.md"
        )

        # Section B is ~5040 chars → should produce multiple sub-chunks
        section_b_chunks = [
            c for c in chunks if "Detailed Regulations" in c["raw_context"]
        ]
        assert len(section_b_chunks) >= 2, (
            f"Expected >= 2 sub-chunks for the massive section, "
            f"got {len(section_b_chunks)}."
        )

        # CRITICAL: Every sub-chunk must still carry the header prefix
        for chunk in section_b_chunks:
            assert "Section:" in chunk["raw_context"], (
                f"Sub-chunk lost its header prefix: {chunk['raw_context'][:100]}…"
            )
            assert "Detailed Regulations" in chunk["raw_context"], (
                "Sub-chunk lost 'Detailed Regulations' header context."
            )

    def test_overlap_in_recursive_sub_chunks(self):
        """
        Sub-chunks from RecursiveCharacterTextSplitter should overlap
        by approximately chunk_overlap characters.
        """
        # Build a Markdown doc where one section has unique, identifiable words
        # Each sentence is unique so we can detect overlap
        sentences = [f"Sentence number {i} describes regulation item {i}. " for i in range(200)]
        body = "".join(sentences)  # ~10000 chars
        md = f"# Doc\n\n## Big Section\n\n{body}\n\n## End\n\nDone.\n"

        parser = MarkdownParser(chunk_size=2000, chunk_overlap=200)
        chunks = parser.parse_text(md, source_file="overlap_test.md")

        # Filter to sub-chunks of the big section that carry sentence content
        # (the first chunk may be a header-only fragment from strip_headers=False)
        import re
        content_chunks = [
            c for c in chunks
            if "Big Section" in c["raw_context"]
            and re.search(r"number \d+", c["raw_context"])
        ]
        assert len(content_chunks) >= 2, (
            f"Expected >= 2 content-bearing sub-chunks, got {len(content_chunks)}"
        )

        # Verify overlap: consecutive chunks should share sentence IDs
        for i in range(len(content_chunks) - 1):
            ids_a = set(re.findall(r"number (\d+)", content_chunks[i]["raw_context"]))
            ids_b = set(re.findall(r"number (\d+)", content_chunks[i + 1]["raw_context"]))
            shared = ids_a & ids_b
            assert len(shared) > 0, (
                f"No overlap detected between sub-chunks {i} and {i+1}. "
                f"IDs A tail: {sorted(ids_a, key=int)[-5:]}, "
                f"IDs B head: {sorted(ids_b, key=int)[:5]}"
            )

    def test_source_file_propagated(
        self, md_parser: MarkdownParser, simple_markdown: str
    ):
        """Every chunk dict should carry the source_file name."""
        chunks = md_parser.parse_text(simple_markdown, source_file="readme.md")
        for chunk in chunks:
            assert chunk["source_file"] == "readme.md"

    def test_plain_text_without_headers(self, md_parser: MarkdownParser):
        """Plain text (no headers) should produce chunks without breadcrumbs."""
        text = "This is plain body text without any markdown headers at all."
        chunks = md_parser.parse_text(text, source_file="notes.txt")
        assert len(chunks) >= 1
        # No sections → no "Section:" prefix
        assert chunks[0]["raw_context"] == text

    def test_empty_content_returns_empty(self, md_parser: MarkdownParser):
        """Empty or whitespace-only input should return no chunks."""
        assert md_parser.parse_text("", source_file="empty.md") == []
        assert md_parser.parse_text("   \n\n  ", source_file="blank.md") == []

    def test_file_read_from_disk(
        self, md_parser: MarkdownParser, tmp_path: Path, simple_markdown: str
    ):
        """parse_file should read from disk and produce identical results."""
        md_file = tmp_path / "design.md"
        md_file.write_text(simple_markdown)

        chunks_from_file = md_parser.parse_file(str(md_file))
        chunks_from_text = md_parser.parse_text(
            simple_markdown, source_file="design.md"
        )

        assert len(chunks_from_file) == len(chunks_from_text)
        # raw_context content should match
        for a, b in zip(chunks_from_file, chunks_from_text):
            assert a["raw_context"] == b["raw_context"]


# ═══════════════════════════════════════════════════════════════════
# 2. WordParser Tests
# ═══════════════════════════════════════════════════════════════════


class TestWordParser:
    """Validate DOCX → Markdown → semantic chunk pipeline."""

    def test_docx_produces_chunks(
        self, word_parser: WordParser, sample_docx: Path
    ):
        """A well-structured DOCX should produce multiple chunks."""
        chunks = word_parser.parse_file(str(sample_docx))
        assert len(chunks) >= 3, (
            f"Expected >= 3 chunks from structured DOCX, got {len(chunks)}."
        )

    def test_docx_headings_appear_in_raw_context(
        self, word_parser: WordParser, sample_docx: Path
    ):
        """
        Word heading styles converted to Markdown headers should
        appear as breadcrumb prefixes in the raw_context.
        """
        chunks = word_parser.parse_file(str(sample_docx))

        # Find chunk about "Credit Line Increase"
        cli_chunks = [
            c for c in chunks if "Credit Line Increase" in c["raw_context"]
        ]
        assert len(cli_chunks) >= 1

        ctx = cli_chunks[0]["raw_context"]
        assert "Section:" in ctx, "DOCX heading was not converted to breadcrumb."

    def test_docx_heading_hierarchy_preserved(
        self, word_parser: WordParser, sample_docx: Path
    ):
        """
        Heading 3 under Heading 2 under Heading 1 should produce a
        multi-level breadcrumb path.
        """
        chunks = word_parser.parse_file(str(sample_docx))

        # Heading 3 "Credit Line Increase" is under H2 "Core Use Cases" under H1 "Business Architecture"
        cli_chunks = [
            c for c in chunks if "Credit Line Increase" in c["raw_context"]
        ]
        assert len(cli_chunks) >= 1
        ctx = cli_chunks[0]["raw_context"]
        assert ">" in ctx, "Expected multi-level breadcrumb with '>'"

    def test_docx_massive_section_triggers_fallback(
        self, word_parser: WordParser, docx_with_massive_section: Path
    ):
        """
        A DOCX heading section exceeding chunk_size should be sub-split
        with headers preserved on every sub-chunk.
        """
        parser = WordParser(chunk_size=2000, chunk_overlap=200)
        chunks = parser.parse_file(str(docx_with_massive_section))

        regulation_chunks = [
            c for c in chunks if "Detailed Regulations" in c["raw_context"]
        ]
        assert len(regulation_chunks) >= 2, (
            f"Expected >= 2 sub-chunks for massive DOCX section, "
            f"got {len(regulation_chunks)}."
        )

        # Every sub-chunk must retain the header
        for chunk in regulation_chunks:
            assert "Section:" in chunk["raw_context"]

    def test_docx_source_file_name(
        self, word_parser: WordParser, sample_docx: Path
    ):
        """Chunks should carry the .docx filename."""
        chunks = word_parser.parse_file(str(sample_docx))
        for chunk in chunks:
            assert chunk["source_file"] == "test_architecture.docx"


# ═══════════════════════════════════════════════════════════════════
# 3. WordParser Internal Conversion Test
# ═══════════════════════════════════════════════════════════════════


class TestWordToMarkdownConversion:
    """Directly test the DOCX → Markdown string conversion."""

    def test_heading_styles_become_markdown_headers(self, sample_docx: Path):
        """Heading 1/2/3 styles should map to #/##/### in the Markdown output."""
        doc = Document(str(sample_docx))
        md_text = WordParser._docx_to_markdown(doc)

        assert "# Business Architecture" in md_text
        assert "## Core Use Cases" in md_text
        assert "### Credit Line Increase" in md_text
        assert "### Automated Risk Mitigation" in md_text
        assert "## Compliance" in md_text

    def test_body_paragraphs_preserved(self, sample_docx: Path):
        """Non-heading paragraphs should appear as plain text."""
        doc = Document(str(sample_docx))
        md_text = WordParser._docx_to_markdown(doc)

        assert "central brain" in md_text
        assert "trace ID" in md_text

    def test_conversion_produces_valid_markdown(self, sample_docx: Path):
        """The output string should be parseable by MarkdownParser without errors."""
        doc = Document(str(sample_docx))
        md_text = WordParser._docx_to_markdown(doc)

        parser = MarkdownParser()
        chunks = parser.parse_text(md_text, source_file="converted.md")
        assert len(chunks) >= 1


# ═══════════════════════════════════════════════════════════════════
# 4. Router Integration Tests
# ═══════════════════════════════════════════════════════════════════


class TestSemanticRouterPhase2:
    """Verify the router now produces real chunks for .md and .docx files."""

    def test_router_processes_markdown_file(
        self, tmp_path: Path, simple_markdown: str
    ):
        """Router should produce ChunkRecords from a .md file."""
        md_file = tmp_path / "design.md"
        md_file.write_text(simple_markdown)

        output_dir = tmp_path / "output"
        config = EmitterConfig(
            output_dir=str(output_dir),
            usecase_id="test-uc",
        )

        with SemanticRouter(config) as router:
            chunks = router.ingest_file(str(md_file))

        assert len(chunks) >= 4
        assert all(isinstance(c, ChunkRecord) for c in chunks)
        assert all(c.usecase_id == "test-uc" for c in chunks)
        assert all(c.file_name == "design.md" for c in chunks)

    def test_router_processes_docx_file(
        self, tmp_path: Path, sample_docx: Path
    ):
        """Router should produce ChunkRecords from a .docx file."""
        output_dir = tmp_path / "output"
        config = EmitterConfig(
            output_dir=str(output_dir),
            usecase_id="credit-uc",
            identifier="underwriting",
        )

        with SemanticRouter(config) as router:
            chunks = router.ingest_file(str(sample_docx))

        assert len(chunks) >= 3
        assert all(c.usecase_id == "credit-uc" for c in chunks)
        assert all(c.identifier == "underwriting" for c in chunks)
        assert all(c.file_name == "test_architecture.docx" for c in chunks)

    def test_router_writes_valid_jsonl(
        self, tmp_path: Path, simple_markdown: str
    ):
        """Router output .jsonl files should contain valid JSON lines."""
        md_file = tmp_path / "test.md"
        md_file.write_text(simple_markdown)

        output_dir = tmp_path / "output"
        config = EmitterConfig(output_dir=str(output_dir))

        with SemanticRouter(config) as router:
            router.ingest_file(str(md_file))

        jsonl_files = list(output_dir.glob("*.jsonl"))
        assert len(jsonl_files) >= 1

        for f in jsonl_files:
            for line in f.read_text().strip().splitlines():
                obj = json.loads(line)
                assert "raw_context" in obj
                assert "Section:" in obj["raw_context"] or "Content:" in obj["raw_context"] or len(obj["raw_context"]) > 0

    def test_router_summary_reflects_processed_files(
        self, tmp_path: Path, simple_markdown: str, sample_docx: Path
    ):
        """Summary should list .md and .docx as processed, not skipped."""
        md_file = tmp_path / "test.md"
        md_file.write_text(simple_markdown)

        output_dir = tmp_path / "output"
        config = EmitterConfig(output_dir=str(output_dir))

        with SemanticRouter(config) as router:
            router.ingest_file(str(md_file))
            router.ingest_file(str(sample_docx))
            summary = router.summary()

        assert len(summary["processed_files"]) == 2
        assert any("test.md" in f for f in summary["processed_files"])
        assert any("test_architecture.docx" in f for f in summary["processed_files"])

    def test_router_yaml_now_processed(self, tmp_path: Path):
        """.yaml files should now be processed by StructuredParser."""
        yaml_file = tmp_path / "config.yaml"
        yaml_file.write_text("database:\n  host: localhost\n  port: 5432\n")

        output_dir = tmp_path / "output"
        config = EmitterConfig(output_dir=str(output_dir))

        with SemanticRouter(config) as router:
            chunks = router.ingest_file(str(yaml_file))

        # Generic YAML → should produce at least 1 chunk via fallback
        assert len(chunks) >= 1

    def test_router_plain_text_uses_markdown_parser(self, tmp_path: Path):
        """Plain .txt files should be routed through MarkdownParser."""
        txt_file = tmp_path / "notes.txt"
        txt_file.write_text("These are plain text notes without any headers.")

        output_dir = tmp_path / "output"
        config = EmitterConfig(output_dir=str(output_dir))

        with SemanticRouter(config) as router:
            chunks = router.ingest_file(str(txt_file))

        assert len(chunks) >= 1
        assert "These are plain text notes without any headers." in chunks[0].raw_context


# ═══════════════════════════════════════════════════════════════════
# 5. Edge Cases
# ═══════════════════════════════════════════════════════════════════


class TestEdgeCases:
    """Edge cases and boundary conditions for the parsers."""

    def test_markdown_with_only_headers_no_body(self, md_parser: MarkdownParser):
        """Headers with no body content should not produce empty chunks."""
        md = "# Title\n\n## Empty Section\n\n## Another Empty\n"
        chunks = md_parser.parse_text(md, source_file="headers_only.md")
        # All chunks should have non-empty raw_context
        for c in chunks:
            assert len(c["raw_context"].strip()) > 0

    def test_markdown_with_code_blocks(self, md_parser: MarkdownParser):
        """Code blocks inside sections should be preserved in raw_context."""
        md = """\
# Setup Guide

## Installation

Run the following command:

```bash
pip install semantic-pipeline
```

This will install all dependencies.
"""
        chunks = md_parser.parse_text(md, source_file="setup.md")
        # Find the installation chunk
        install_chunks = [c for c in chunks if "Installation" in c["raw_context"]]
        assert len(install_chunks) >= 1
        assert "pip install" in install_chunks[0]["raw_context"]

    def test_docx_empty_paragraphs_ignored(self, tmp_path: Path):
        """DOCX with empty paragraphs should not produce empty chunks."""
        doc = Document()
        doc.add_heading("Title", level=1)
        doc.add_paragraph("")
        doc.add_paragraph("")
        doc.add_paragraph("Actual content here.")

        path = tmp_path / "sparse.docx"
        doc.save(str(path))

        parser = WordParser()
        chunks = parser.parse_file(str(path))
        for c in chunks:
            content = c["raw_context"].strip()
            assert len(content) > 0
