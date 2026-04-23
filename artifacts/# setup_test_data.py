import json
import yaml
import os
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
import matplotlib.pyplot as plt
from PIL import Image, ImageDraw, ImageFont

def generate_json_image(filename, json_data, title):
    """Utility to turn a JSON dict into an image to embed in Excel."""
    json_str = json.dumps(json_data, indent=2)
    # create figure
    fig, ax = plt.subplots(figsize=(6, 4))
    ax.axis('off')
    ax.text(0.01, 0.99, f"{title}\n\n{json_str}", 
            fontsize=9, family='monospace', va='top', ha='left')
    plt.savefig(filename, bbox_inches='tight', dpi=150)
    plt.close()

def generate_openapi():
    openapi_spec = {
        "openapi": "3.0.3",
        "info": {
            "title": "Credit Decisioning Microservice API",
            "description": "API for underwriting and credit line increase decisions. Includes rule evaluation, manual review queues, and customer profile integration.",
            "version": "2.5.0",
            "contact": {
                "name": "Underwriting Tech Team",
                "email": "underwriting-tech@bank.com"
            }
        },
        "servers": [
            {"url": "https://api.bank.com/v1/credit-decisioning", "description": "Production Server"},
            {"url": "https://staging-api.bank.com/v1/credit-decisioning", "description": "Staging Server"}
        ],
        "tags": [
            {"name": "Application", "description": "Endpoints managing credit line applications"},
            {"name": "Decision", "description": "Core decision engine and evaluation logic"},
            {"name": "CustomerProfile", "description": "Customer financial profile and bureau data"},
            {"name": "RuleEngine", "description": "Underwriting rules management"},
            {"name": "Reporting", "description": "Metrics and audit reporting"}
        ],
        "paths": {}
    }
    
    # We will procedurally generate endpoints to make the file large and comprehensive.
    
    # 1. ApplicationController (6 endpoints)
    application_paths = {
        "/applications": {
            "post": {"tags": ["Application"], "summary": "Submit a new application", "responses": {"201": {"description": "Application created"}}},
            "get": {"tags": ["Application"], "summary": "List all applications", "responses": {"200": {"description": "OK"}}}
        },
        "/applications/{appId}": {
            "get": {"tags": ["Application"], "summary": "Get application status", "responses": {"200": {"description": "OK"}}},
            "patch": {"tags": ["Application"], "summary": "Update application details", "responses": {"200": {"description": "OK"}}},
            "delete": {"tags": ["Application"], "summary": "Withdraw application", "responses": {"204": {"description": "Withdrawn"}}}
        },
        "/applications/{appId}/documents": {
            "post": {"tags": ["Application"], "summary": "Upload supporting documents (paystubs etc)", "responses": {"202": {"description": "Uploaded"}}}
        }
    }
    
    # 2. DecisionController (8 endpoints)
    decision_paths = {
        "/decisions/evaluate": {
            "post": {"tags": ["Decision"], "summary": "Evaluate an application automatically", "responses": {"200": {"description": "Decision rendered (Approve/Decline/Manual)"}}}
        },
        "/decisions/manual-review": {
            "get": {"tags": ["Decision"], "summary": "Get applications pending manual review", "responses": {"200": {"description": "Queue list"}}},
            "post": {"tags": ["Decision"], "summary": "Submit manual review decision", "responses": {"200": {"description": "Decision updated"}}}
        },
        "/decisions/{decisionId}/override": {
            "post": {"tags": ["Decision"], "summary": "Override a system decision (Admin only)", "responses": {"200": {"description": "Overridden"}}}
        },
        "/decisions/{decisionId}/approve": {
            "post": {"tags": ["Decision"], "summary": "Force approve an application", "responses": {"200": {"description": "Approved"}}}
        },
        "/decisions/{decisionId}/decline": {
            "post": {"tags": ["Decision"], "summary": "Force decline an application", "responses": {"200": {"description": "Declined"}}}
        },
        "/decisions/{decisionId}/request-info": {
            "post": {"tags": ["Decision"], "summary": "Request more info from applicant", "responses": {"200": {"description": "Status changed to Pending Info"}}}
        },
        "/decisions/history/{customerId}": {
            "get": {"tags": ["Decision"], "summary": "Get decision history for a customer", "responses": {"200": {"description": "History logs"}}}
        }
    }

    # 3. CustomerProfileController (6 endpoints)
    profile_paths = {
        "/profiles/{customerId}": {
            "get": {"tags": ["CustomerProfile"], "summary": "Get customer profile", "responses": {"200": {"description": "Profile data"}}},
            "put": {"tags": ["CustomerProfile"], "summary": "Update profile", "responses": {"200": {"description": "Updated"}}}
        },
        "/profiles/{customerId}/income": {
            "post": {"tags": ["CustomerProfile"], "summary": "Add verified income", "responses": {"200": {"description": "Income added"}}}
        },
        "/profiles/{customerId}/identity": {
            "post": {"tags": ["CustomerProfile"], "summary": "Trigger KYC/Identity verification", "responses": {"200": {"description": "Verified"}}}
        },
        "/profiles/{customerId}/bureau": {
            "post": {"tags": ["CustomerProfile"], "summary": "Pull latest credit bureau report", "responses": {"200": {"description": "Report pulled"}}},
            "get": {"tags": ["CustomerProfile"], "summary": "Get cached credit score", "responses": {"200": {"description": "Score details"}}}
        }
    }
    
    # 4. RuleEngineController (6 endpoints)
    rule_paths = {
        "/rules": {
            "get": {"tags": ["RuleEngine"], "summary": "List all active underwriting rules", "responses": {"200": {"description": "Rules list"}}},
            "post": {"tags": ["RuleEngine"], "summary": "Create a new rule", "responses": {"201": {"description": "Rule created"}}}
        },
        "/rules/{ruleId}": {
            "get": {"tags": ["RuleEngine"], "summary": "Get rule definition", "responses": {"200": {"description": "Rule data"}}},
            "put": {"tags": ["RuleEngine"], "summary": "Update rule expression", "responses": {"200": {"description": "Rule updated"}}},
            "delete": {"tags": ["RuleEngine"], "summary": "Disable a rule", "responses": {"204": {"description": "Disabled"}}}
        },
        "/rules/simulate": {
            "post": {"tags": ["RuleEngine"], "summary": "Test a rule against mock data without saving", "responses": {"200": {"description": "Simulation result"}}}
        }
    }

    # 5. ReportingController (6 endpoints)
    reporting_paths = {
        "/reports/metrics": {
            "get": {"tags": ["Reporting"], "summary": "Get high-level decision metrics (approval rates)", "responses": {"200": {"description": "Metrics"}}}
        },
        "/reports/export": {
            "get": {"tags": ["Reporting"], "summary": "Export decisions to CSV", "responses": {"200": {"description": "CSV stream"}}}
        },
        "/reports/audit": {
            "get": {"tags": ["Reporting"], "summary": "Get compliance audit log", "responses": {"200": {"description": "Audit paginated"}}}
        },
        "/reports/kpi": {
            "get": {"tags": ["Reporting"], "summary": "Get KPI dashboard widget data", "responses": {"200": {"description": "Dashboard JSON"}}}
        },
        "/reports/sla": {
            "get": {"tags": ["Reporting"], "summary": "System SLA and latency stats", "responses": {"200": {"description": "Wait times"}}}
        },
        "/reports/fair-lending": {
            "get": {"tags": ["Reporting"], "summary": "Fair lending compliance check report", "responses": {"200": {"description": "Compliance metrics"}}}
        }
    }
    
    # Merge all
    openapi_spec["paths"].update(application_paths)
    openapi_spec["paths"].update(decision_paths)
    openapi_spec["paths"].update(profile_paths)
    openapi_spec["paths"].update(rule_paths)
    openapi_spec["paths"].update(reporting_paths)

    with open('credit_decisioning_openapi.yaml', 'w') as f:
        yaml.dump(openapi_spec, f, sort_keys=False)
        
    with open('credit_decisioning_swagger.json', 'w') as f:
        json.dump(openapi_spec, f, indent=2)
        
    print("Created Swagger JSON and OpenAPI YAML.")

def generate_word_doc():
    doc = Document()
    
    # Title
    title = doc.add_heading('Business Architecture: Credit Decisioning Microservice', 0)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    
    # Executive Summary
    doc.add_heading('1. Executive Summary', level=1)
    doc.add_paragraph('The Credit Decisioning Microservice acts as the central brain '
                      'for the automated underwriting and credit line increase (CLI) workflows. '
                      'Its primary purpose is to assess customer risk profiles instantaneously '
                      'and return a mathematically sound decision on whether to extend, maintain, '
                      'or reduce credit limits.')
    
    doc.add_paragraph('By migrating to a dedicated microservice, the business aims to achieve a '
                      'straight-through processing (STP) rate of over 85%, significantly reducing '
                      'manual underwriting costs while adhering strictly to Fair Lending regulations.')
                      
    # Business Use Cases
    doc.add_heading('2. Core Business Use Cases', level=1)
    doc.add_heading('2.1 Credit Line Increase (CLI) Origination', level=2)
    doc.add_paragraph('Customers can request a credit line increase via the mobile app or web portal. '
                      'This triggers a real-time call to the Decisioning Engine. The engine pulls '
                      'data from internal payment histories, external bureau files, and income verification '
                      'modules to approve or decline the CLI instantly.')
                      
    doc.add_heading('2.2 Automated Risk Mitigation', level=2)
    doc.add_paragraph('During economic downturns, the microservice actively runs batch rules against '
                      'high-risk portfolios to proactively decrease exposure before defaults occur. '
                      'It acts as an autonomous risk-adjustment layer.')
                      
    # Process Funnel
    doc.add_heading('3. The Underwriting Funnel', level=1)
    doc.add_paragraph('1. Data Aggregation: Internal and external APIs supply the applicant context.\n'
                      '2. Knockout Rules: Immediate failure criteria (e.g. bankruptcies in last 12 months, age < 18).\n'
                      '3. Scoring: ML model applies a custom credit score (Internal Risk Score).\n'
                      '4. Limit Assignment: Based on risk brackets, the system assigns a dynamic credit ceiling.\n'
                      '5. Manual Review Fallback: Marginal cases drop into a queue for human underwriters.')

    doc.add_heading('4. Compliance and Auditability', level=1)
    doc.add_paragraph('Every decision yields a unique trace ID linking the exact rules, '
                      'bureau scores, and timestamps used. This guarantees complete explainability '
                      'for adverse action notices (FCRA compliance).')

    doc.save('business_architecture.docx')
    print("Created Word Document (Business Architecture).")

def generate_markdown():
    md_content = """# High-Level Design (HLD): Credit Decisioning Microservice

## 1. System Overview
The Decisioning Microservice evaluates user applications for Credit Line Increases (CLI) and auto-generates underwriting decisions. It ensures sub-second latency for UI-facing calls while maintaining eventual consistency across reporting engines.

## 2. Architecture Diagram (Mermaid)
```mermaid
graph TD
    A[Mobile/Web Client] -->|API Gateway| B(Credit Decisioning API)
    B --> C{Rule Engine}
    B --> D[Customer Profile Cache]
    C -->|Fetch Bureau| E[External Credit Bureau API]
    C -->|Fetch Income| F[Internal Core Banking DB]
    C -->|Publish Event| G((Kafka Topic: decisions))
    G --> H[Reporting/Audit Service]
    G --> I[Manual Review UI]
```

## 3. Technology Stack
- **Language**: Java / Spring Boot 3
- **Database**: PostgreSQL (Relational schema for configurations and audit logs)
- **Cache**: Redis (For caching bureau scores up to 30 days)
- **Messaging**: Apache Kafka
- **Infrastructure**: Kubernetes (EKS), highly available multi-AZ deployment.

## 4. Key Components
### 4.1 Rule Engine
Evaluates deterministic logic trees configured by business analysts (e.g., `IF income > 50000 AND DTI < 40% THEN APPROVE`).

### 4.2 Data Aggregator
Asynchronously fetches data from external bureaus (Experian/Equifax) and internal ledgers.

### 4.3 Policy Router
Determines the routing of the decision: Auto-Approve, Auto-Decline, or Manual Review.
"""
    with open('high_level_design.md', 'w') as f:
        f.write(md_content)
    print("Created Markdown HLD.")

def generate_excel():
    excel_file = 'decision_controller_rules.xlsx'
    
    # Define the 8 endpoints for the DecisionController
    endpoints = [
        {"name": "Evaluate App", "path": "/decisions/evaluate", "method": "POST", "desc": "Evaluates a CLI application via rules engine."},
        {"name": "Manual Review Queue", "path": "/decisions/manual-review", "method": "GET", "desc": "Fetches apps routed to human underwriters."},
        {"name": "Submit Manual Dec", "path": "/decisions/manual-review", "method": "POST", "desc": "Human underwriter commits an override."},
        {"name": "Admin Override", "path": "/decisions/{decisionId}/override", "method": "POST", "desc": "Admin ignores rules and forces a state."},
        {"name": "Force Approve", "path": "/decisions/{decisionId}/approve", "method": "POST", "desc": "Bypass logic to grant line increase."},
        {"name": "Force Decline", "path": "/decisions/{decisionId}/decline", "method": "POST", "desc": "Bypass logic to deny line increase."},
        {"name": "Request Info", "path": "/decisions/{decisionId}/request-info", "method": "POST", "desc": "Marks app as pending, notifies user for documents."},
        {"name": "Decision History", "path": "/decisions/history/{customerId}", "method": "GET", "desc": "Retrieves the historical log of all decisions for a user."}
    ]
    
    with pd.ExcelWriter(excel_file, engine='xlsxwriter') as writer:
        workbook = writer.book
        
        # Add formatting
        header_format = workbook.add_format({'bold': True, 'bg_color': '#4F81BD', 'font_color': 'white', 'border': 1})
        cell_format = workbook.add_format({'text_wrap': True, 'valign': 'top'})
        
        for ep in endpoints:
            # Sheet names max 31 chars
            safe_sheet_name = ep['name'][:31].replace('/', '_')
            worksheet = workbook.add_worksheet(safe_sheet_name)
            
            # Write Header details
            worksheet.write(0, 0, "Endpoint Name:", header_format)
            worksheet.write(0, 1, ep['name'])
            worksheet.write(1, 0, "Method & Path:", header_format)
            worksheet.write(1, 1, f"[{ep['method']}] {ep['path']}")
            worksheet.write(2, 0, "Description:", header_format)
            worksheet.write(2, 1, ep['desc'])
            
            # Table Structure for Fields and rules
            headers = ["Field Key", "Type", "Required", "Business Rule/Logic", "Sample Value"]
            for col_num, header_title in enumerate(headers):
                worksheet.write(4, col_num, header_title, header_format)
                
            # Dummy rules based on the endpoint
            if 'evaluate' in ep['path']:
                rows = [
                    ["customerId", "UUID", "Y", "Lookup core profile", "123e4567-e89b-12d3"],
                    ["requestedAmount", "Decimal", "Y", "Must be > current Limit", "5000.00"],
                    ["statedIncome", "Decimal", "Y", "Must evaluate DTI < 40%", "120000.00"]
                ]
                req = {"customerId": "123e4567-e89b-12d3", "requestedAmount": 5000, "statedIncome": 120000}
                res = {"decision": "APPROVED", "approvedAmount": 5000, "reasonCode": "A01"}
            elif 'manual' in ep['path'] and ep['method'] == 'POST':
                rows = [
                    ["underwriterId", "String", "Y", "Must have correct RBAC role", "U-9942"],
                    ["decision", "String", "Y", "ENUM: APPROVE, DECLINE", "APPROVE"],
                    ["notes", "String", "N", "Audit trail reasoning", "Verified paystub looks good"]
                ]
                req = {"underwriterId": "U-9942", "decision": "APPROVE", "notes": "Everything ok"}
                res = {"status": "SUCCESS", "updatedAt": "2026-04-23T10:00:00Z"}
            elif 'manual' in ep['path'] and ep['method'] == 'GET':
                 rows = [
                    ["status", "String", "N", "Filter by status: PENDING", "PENDING"],
                    ["limit", "Integer", "N", "Pagination logic", "100"]
                 ]
                 req = {} # Query params
                 res = {"page": 1, "items": [{"decisionId": "d-123", "score": 680, "status": "PENDING"}]}
            else:
                rows = [
                    ["decisionId", "UUID", "Y", "Must be valid in DB", "d-84848"],
                    ["reason", "String", "Y", "Required for compliance logs", "Customer requested"]
                ]
                req = {"reason": "Customer requested account closure"}
                res = {"status": "ACKNOWLEDGED"}
                
            for row_num, row_data in enumerate(rows, start=5):
                for col_num, cell_data in enumerate(row_data):
                    worksheet.write(row_num, col_num, cell_data, cell_format)
                    
            worksheet.set_column('A:A', 20)
            worksheet.set_column('B:E', 30)
            
            # Generate and insert sample images
            req_img = f"req_{safe_sheet_name}.png"
            res_img = f"res_{safe_sheet_name}.png"
            generate_json_image(req_img, req, "SAMPLE REQUEST PAYLOAD")
            generate_json_image(res_img, res, "SAMPLE RESPONSE PAYLOAD")
            
            worksheet.write(10, 0, "Sample Request:", header_format)
            worksheet.insert_image(11, 0, req_img)
            worksheet.write(10, 4, "Sample Response:", header_format)
            worksheet.insert_image(11, 4, res_img)
            
    print("Created Excel file mapped to endpoints with Request/Response schemas embedded.")
    
    # Cleanup temporary images
    for ep in endpoints:
        safe_sheet_name = ep['name'][:31].replace('/', '_')
        if os.path.exists(f"req_{safe_sheet_name}.png"): os.remove(f"req_{safe_sheet_name}.png")
        if os.path.exists(f"res_{safe_sheet_name}.png"): os.remove(f"res_{safe_sheet_name}.png")

if __name__ == "__main__":
    generate_openapi()
    generate_word_doc()
    generate_markdown()
    generate_excel()
    print("All comprehensive artifacts correctly generated.")import json
import yaml
import os
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
import matplotlib.pyplot as plt
from PIL import Image, ImageDraw, ImageFont

def generate_json_image(filename, json_data, title):
    """Utility to turn a JSON dict into an image to embed in Excel."""
    json_str = json.dumps(json_data, indent=2)
    # create figure
    fig, ax = plt.subplots(figsize=(6, 4))
    ax.axis('off')
    ax.text(0.01, 0.99, f"{title}\n\n{json_str}", 
            fontsize=9, family='monospace', va='top', ha='left')
    plt.savefig(filename, bbox_inches='tight', dpi=150)
    plt.close()

def generate_openapi():
    openapi_spec = {
        "openapi": "3.0.3",
        "info": {
            "title": "Credit Decisioning Microservice API",
            "description": "API for underwriting and credit line increase decisions. Includes rule evaluation, manual review queues, and customer profile integration.",
            "version": "2.5.0",
            "contact": {
                "name": "Underwriting Tech Team",
                "email": "underwriting-tech@bank.com"
            }
        },
        "servers": [
            {"url": "https://api.bank.com/v1/credit-decisioning", "description": "Production Server"},
            {"url": "https://staging-api.bank.com/v1/credit-decisioning", "description": "Staging Server"}
        ],
        "tags": [
            {"name": "Application", "description": "Endpoints managing credit line applications"},
            {"name": "Decision", "description": "Core decision engine and evaluation logic"},
            {"name": "CustomerProfile", "description": "Customer financial profile and bureau data"},
            {"name": "RuleEngine", "description": "Underwriting rules management"},
            {"name": "Reporting", "description": "Metrics and audit reporting"}
        ],
        "paths": {}
    }
    
    # We will procedurally generate endpoints to make the file large and comprehensive.
    
    # 1. ApplicationController (6 endpoints)
    application_paths = {
        "/applications": {
            "post": {"tags": ["Application"], "summary": "Submit a new application", "responses": {"201": {"description": "Application created"}}},
            "get": {"tags": ["Application"], "summary": "List all applications", "responses": {"200": {"description": "OK"}}}
        },
        "/applications/{appId}": {
            "get": {"tags": ["Application"], "summary": "Get application status", "responses": {"200": {"description": "OK"}}},
            "patch": {"tags": ["Application"], "summary": "Update application details", "responses": {"200": {"description": "OK"}}},
            "delete": {"tags": ["Application"], "summary": "Withdraw application", "responses": {"204": {"description": "Withdrawn"}}}
        },
        "/applications/{appId}/documents": {
            "post": {"tags": ["Application"], "summary": "Upload supporting documents (paystubs etc)", "responses": {"202": {"description": "Uploaded"}}}
        }
    }
    
    # 2. DecisionController (8 endpoints)
    decision_paths = {
        "/decisions/evaluate": {
            "post": {"tags": ["Decision"], "summary": "Evaluate an application automatically", "responses": {"200": {"description": "Decision rendered (Approve/Decline/Manual)"}}}
        },
        "/decisions/manual-review": {
            "get": {"tags": ["Decision"], "summary": "Get applications pending manual review", "responses": {"200": {"description": "Queue list"}}},
            "post": {"tags": ["Decision"], "summary": "Submit manual review decision", "responses": {"200": {"description": "Decision updated"}}}
        },
        "/decisions/{decisionId}/override": {
            "post": {"tags": ["Decision"], "summary": "Override a system decision (Admin only)", "responses": {"200": {"description": "Overridden"}}}
        },
        "/decisions/{decisionId}/approve": {
            "post": {"tags": ["Decision"], "summary": "Force approve an application", "responses": {"200": {"description": "Approved"}}}
        },
        "/decisions/{decisionId}/decline": {
            "post": {"tags": ["Decision"], "summary": "Force decline an application", "responses": {"200": {"description": "Declined"}}}
        },
        "/decisions/{decisionId}/request-info": {
            "post": {"tags": ["Decision"], "summary": "Request more info from applicant", "responses": {"200": {"description": "Status changed to Pending Info"}}}
        },
        "/decisions/history/{customerId}": {
            "get": {"tags": ["Decision"], "summary": "Get decision history for a customer", "responses": {"200": {"description": "History logs"}}}
        }
    }

    # 3. CustomerProfileController (6 endpoints)
    profile_paths = {
        "/profiles/{customerId}": {
            "get": {"tags": ["CustomerProfile"], "summary": "Get customer profile", "responses": {"200": {"description": "Profile data"}}},
            "put": {"tags": ["CustomerProfile"], "summary": "Update profile", "responses": {"200": {"description": "Updated"}}}
        },
        "/profiles/{customerId}/income": {
            "post": {"tags": ["CustomerProfile"], "summary": "Add verified income", "responses": {"200": {"description": "Income added"}}}
        },
        "/profiles/{customerId}/identity": {
            "post": {"tags": ["CustomerProfile"], "summary": "Trigger KYC/Identity verification", "responses": {"200": {"description": "Verified"}}}
        },
        "/profiles/{customerId}/bureau": {
            "post": {"tags": ["CustomerProfile"], "summary": "Pull latest credit bureau report", "responses": {"200": {"description": "Report pulled"}}},
            "get": {"tags": ["CustomerProfile"], "summary": "Get cached credit score", "responses": {"200": {"description": "Score details"}}}
        }
    }
    
    # 4. RuleEngineController (6 endpoints)
    rule_paths = {
        "/rules": {
            "get": {"tags": ["RuleEngine"], "summary": "List all active underwriting rules", "responses": {"200": {"description": "Rules list"}}},
            "post": {"tags": ["RuleEngine"], "summary": "Create a new rule", "responses": {"201": {"description": "Rule created"}}}
        },
        "/rules/{ruleId}": {
            "get": {"tags": ["RuleEngine"], "summary": "Get rule definition", "responses": {"200": {"description": "Rule data"}}},
            "put": {"tags": ["RuleEngine"], "summary": "Update rule expression", "responses": {"200": {"description": "Rule updated"}}},
            "delete": {"tags": ["RuleEngine"], "summary": "Disable a rule", "responses": {"204": {"description": "Disabled"}}}
        },
        "/rules/simulate": {
            "post": {"tags": ["RuleEngine"], "summary": "Test a rule against mock data without saving", "responses": {"200": {"description": "Simulation result"}}}
        }
    }

    # 5. ReportingController (6 endpoints)
    reporting_paths = {
        "/reports/metrics": {
            "get": {"tags": ["Reporting"], "summary": "Get high-level decision metrics (approval rates)", "responses": {"200": {"description": "Metrics"}}}
        },
        "/reports/export": {
            "get": {"tags": ["Reporting"], "summary": "Export decisions to CSV", "responses": {"200": {"description": "CSV stream"}}}
        },
        "/reports/audit": {
            "get": {"tags": ["Reporting"], "summary": "Get compliance audit log", "responses": {"200": {"description": "Audit paginated"}}}
        },
        "/reports/kpi": {
            "get": {"tags": ["Reporting"], "summary": "Get KPI dashboard widget data", "responses": {"200": {"description": "Dashboard JSON"}}}
        },
        "/reports/sla": {
            "get": {"tags": ["Reporting"], "summary": "System SLA and latency stats", "responses": {"200": {"description": "Wait times"}}}
        },
        "/reports/fair-lending": {
            "get": {"tags": ["Reporting"], "summary": "Fair lending compliance check report", "responses": {"200": {"description": "Compliance metrics"}}}
        }
    }
    
    # Merge all
    openapi_spec["paths"].update(application_paths)
    openapi_spec["paths"].update(decision_paths)
    openapi_spec["paths"].update(profile_paths)
    openapi_spec["paths"].update(rule_paths)
    openapi_spec["paths"].update(reporting_paths)

    with open('credit_decisioning_openapi.yaml', 'w') as f:
        yaml.dump(openapi_spec, f, sort_keys=False)
        
    with open('credit_decisioning_swagger.json', 'w') as f:
        json.dump(openapi_spec, f, indent=2)
        
    print("Created Swagger JSON and OpenAPI YAML.")

def generate_word_doc():
    doc = Document()
    
    # Title
    title = doc.add_heading('Business Architecture: Credit Decisioning Microservice', 0)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    
    # Executive Summary
    doc.add_heading('1. Executive Summary', level=1)
    doc.add_paragraph('The Credit Decisioning Microservice acts as the central brain '
                      'for the automated underwriting and credit line increase (CLI) workflows. '
                      'Its primary purpose is to assess customer risk profiles instantaneously '
                      'and return a mathematically sound decision on whether to extend, maintain, '
                      'or reduce credit limits.')
    
    doc.add_paragraph('By migrating to a dedicated microservice, the business aims to achieve a '
                      'straight-through processing (STP) rate of over 85%, significantly reducing '
                      'manual underwriting costs while adhering strictly to Fair Lending regulations.')
                      
    # Business Use Cases
    doc.add_heading('2. Core Business Use Cases', level=1)
    doc.add_heading('2.1 Credit Line Increase (CLI) Origination', level=2)
    doc.add_paragraph('Customers can request a credit line increase via the mobile app or web portal. '
                      'This triggers a real-time call to the Decisioning Engine. The engine pulls '
                      'data from internal payment histories, external bureau files, and income verification '
                      'modules to approve or decline the CLI instantly.')
                      
    doc.add_heading('2.2 Automated Risk Mitigation', level=2)
    doc.add_paragraph('During economic downturns, the microservice actively runs batch rules against '
                      'high-risk portfolios to proactively decrease exposure before defaults occur. '
                      'It acts as an autonomous risk-adjustment layer.')
                      
    # Process Funnel
    doc.add_heading('3. The Underwriting Funnel', level=1)
    doc.add_paragraph('1. Data Aggregation: Internal and external APIs supply the applicant context.\n'
                      '2. Knockout Rules: Immediate failure criteria (e.g. bankruptcies in last 12 months, age < 18).\n'
                      '3. Scoring: ML model applies a custom credit score (Internal Risk Score).\n'
                      '4. Limit Assignment: Based on risk brackets, the system assigns a dynamic credit ceiling.\n'
                      '5. Manual Review Fallback: Marginal cases drop into a queue for human underwriters.')

    doc.add_heading('4. Compliance and Auditability', level=1)
    doc.add_paragraph('Every decision yields a unique trace ID linking the exact rules, '
                      'bureau scores, and timestamps used. This guarantees complete explainability '
                      'for adverse action notices (FCRA compliance).')

    doc.save('business_architecture.docx')
    print("Created Word Document (Business Architecture).")

def generate_markdown():
    md_content = """# High-Level Design (HLD): Credit Decisioning Microservice

## 1. System Overview
The Decisioning Microservice evaluates user applications for Credit Line Increases (CLI) and auto-generates underwriting decisions. It ensures sub-second latency for UI-facing calls while maintaining eventual consistency across reporting engines.

## 2. Architecture Diagram (Mermaid)
```mermaid
graph TD
    A[Mobile/Web Client] -->|API Gateway| B(Credit Decisioning API)
    B --> C{Rule Engine}
    B --> D[Customer Profile Cache]
    C -->|Fetch Bureau| E[External Credit Bureau API]
    C -->|Fetch Income| F[Internal Core Banking DB]
    C -->|Publish Event| G((Kafka Topic: decisions))
    G --> H[Reporting/Audit Service]
    G --> I[Manual Review UI]
```

## 3. Technology Stack
- **Language**: Java / Spring Boot 3
- **Database**: PostgreSQL (Relational schema for configurations and audit logs)
- **Cache**: Redis (For caching bureau scores up to 30 days)
- **Messaging**: Apache Kafka
- **Infrastructure**: Kubernetes (EKS), highly available multi-AZ deployment.

## 4. Key Components
### 4.1 Rule Engine
Evaluates deterministic logic trees configured by business analysts (e.g., `IF income > 50000 AND DTI < 40% THEN APPROVE`).

### 4.2 Data Aggregator
Asynchronously fetches data from external bureaus (Experian/Equifax) and internal ledgers.

### 4.3 Policy Router
Determines the routing of the decision: Auto-Approve, Auto-Decline, or Manual Review.
"""
    with open('high_level_design.md', 'w') as f:
        f.write(md_content)
    print("Created Markdown HLD.")

def generate_excel():
    excel_file = 'decision_controller_rules.xlsx'
    
    # Define the 8 endpoints for the DecisionController
    endpoints = [
        {"name": "Evaluate App", "path": "/decisions/evaluate", "method": "POST", "desc": "Evaluates a CLI application via rules engine."},
        {"name": "Manual Review Queue", "path": "/decisions/manual-review", "method": "GET", "desc": "Fetches apps routed to human underwriters."},
        {"name": "Submit Manual Dec", "path": "/decisions/manual-review", "method": "POST", "desc": "Human underwriter commits an override."},
        {"name": "Admin Override", "path": "/decisions/{decisionId}/override", "method": "POST", "desc": "Admin ignores rules and forces a state."},
        {"name": "Force Approve", "path": "/decisions/{decisionId}/approve", "method": "POST", "desc": "Bypass logic to grant line increase."},
        {"name": "Force Decline", "path": "/decisions/{decisionId}/decline", "method": "POST", "desc": "Bypass logic to deny line increase."},
        {"name": "Request Info", "path": "/decisions/{decisionId}/request-info", "method": "POST", "desc": "Marks app as pending, notifies user for documents."},
        {"name": "Decision History", "path": "/decisions/history/{customerId}", "method": "GET", "desc": "Retrieves the historical log of all decisions for a user."}
    ]
    
    with pd.ExcelWriter(excel_file, engine='xlsxwriter') as writer:
        workbook = writer.book
        
        # Add formatting
        header_format = workbook.add_format({'bold': True, 'bg_color': '#4F81BD', 'font_color': 'white', 'border': 1})
        cell_format = workbook.add_format({'text_wrap': True, 'valign': 'top'})
        
        for ep in endpoints:
            # Sheet names max 31 chars
            safe_sheet_name = ep['name'][:31].replace('/', '_')
            worksheet = workbook.add_worksheet(safe_sheet_name)
            
            # Write Header details
            worksheet.write(0, 0, "Endpoint Name:", header_format)
            worksheet.write(0, 1, ep['name'])
            worksheet.write(1, 0, "Method & Path:", header_format)
            worksheet.write(1, 1, f"[{ep['method']}] {ep['path']}")
            worksheet.write(2, 0, "Description:", header_format)
            worksheet.write(2, 1, ep['desc'])
            
            # Table Structure for Fields and rules
            headers = ["Field Key", "Type", "Required", "Business Rule/Logic", "Sample Value"]
            for col_num, header_title in enumerate(headers):
                worksheet.write(4, col_num, header_title, header_format)
                
            # Dummy rules based on the endpoint
            if 'evaluate' in ep['path']:
                rows = [
                    ["customerId", "UUID", "Y", "Lookup core profile", "123e4567-e89b-12d3"],
                    ["requestedAmount", "Decimal", "Y", "Must be > current Limit", "5000.00"],
                    ["statedIncome", "Decimal", "Y", "Must evaluate DTI < 40%", "120000.00"]
                ]
                req = {"customerId": "123e4567-e89b-12d3", "requestedAmount": 5000, "statedIncome": 120000}
                res = {"decision": "APPROVED", "approvedAmount": 5000, "reasonCode": "A01"}
            elif 'manual' in ep['path']:
                rows = [
                    ["underwriterId", "String", "Y", "Must have correct RBAC role", "U-9942"],
                    ["decision", "String", "Y", "ENUM: APPROVE, DECLINE", "APPROVE"],
                    ["notes", "String", "N", "Audit trail reasoning", "Verified paystub looks good"]
                ]
                req = {"underwriterId": "U-9942", "decision": "APPROVE", "notes": "Everything ok"}
                res = {"status": "SUCCESS", "updatedAt": "2026-04-23T10:00:00Z"}
            else:
                rows = [
                    ["decisionId", "UUID", "Y", "Must be valid in DB", "d-84848"],
                    ["reason", "String", "Y", "Required for compliance logs", "Customer requested"]
                ]
                req = {"reason": "Customer requested account closure"}
                res = {"status": "ACKNOWLEDGED"}
                
            for row_num, row_data in enumerate(rows, start=5):
                for col_num, cell_data in enumerate(row_data):
                    worksheet.write(row_num, col_num, cell_data, cell_format)
                    
            worksheet.set_column('A:A', 20)
            worksheet.set_column('B:E', 30)
            
            # Generate and insert sample images
            req_img = f"req_{safe_sheet_name}.png"
            res_img = f"res_{safe_sheet_name}.png"
            generate_json_image(req_img, req, "SAMPLE REQUEST PAYLOAD")
            generate_json_image(res_img, res, "SAMPLE RESPONSE PAYLOAD")
            
            worksheet.write(10, 0, "Sample Request:", header_format)
            worksheet.insert_image(11, 0, req_img)
            worksheet.write(10, 4, "Sample Response:", header_format)
            worksheet.insert_image(11, 4, res_img)
            
    print("Created Excel file mapped to endpoints with Request/Response schemas embedded.")
    
    # Cleanup temporary images
    for ep in endpoints:
        safe_sheet_name = ep['name'][:31].replace('/', '_')
        if os.path.exists(f"req_{safe_sheet_name}.png"): os.remove(f"req_{safe_sheet_name}.png")
        if os.path.exists(f"res_{safe_sheet_name}.png"): os.remove(f"res_{safe_sheet_name}.png")

if __name__ == "__main__":
    generate_openapi()
    generate_word_doc()
    generate_markdown()
    generate_excel()
    print("All comprehensive artifacts correctly generated.")
# setup_test_data.py
import pandas as pd
import numpy as np
import json
import yaml
from docx import Document
from docx.shared import Inches
import os
import random
from faker import Faker
import matplotlib.pyplot as plt

fake = Faker()

def generate_dummy_images():
    # Create simple dummy image
    plt.figure(figsize=(4, 3))
    plt.plot([1, 2, 3], [10, 20, 15])
    plt.title("Sample Chart Image")
    plt.savefig('dummy_chart.png')
    plt.close()

# 1. Generate Complex Excel (10 sheets, lots of data, charts, images)
def create_excel():
    num_sheets = 10
    rows_per_sheet = 5000
    excel_file = 'sample_complex.xlsx'
    
    # Needs dummy image to insert
    generate_dummy_images()
    
    with pd.ExcelWriter(excel_file, engine='xlsxwriter') as writer:
        workbook = writer.book
        
        for i in range(1, num_sheets + 1):
            sheet_name = f'Department_Data_{i}'
            
            # Generate big DataFrame
            data = {
                'ID': range(1, rows_per_sheet + 1),
                'Name': [fake.name() for _ in range(rows_per_sheet)],
                'Job Title': [fake.job() for _ in range(rows_per_sheet)],
                'Salary': np.random.randint(40000, 200000, size=rows_per_sheet),
                'Start Date': [fake.date_between(start_date='-10y', end_date='today') for _ in range(rows_per_sheet)],
                'Performance_Score': np.random.uniform(1.0, 5.0, size=rows_per_sheet),
                'Category': np.random.choice(['A', 'B', 'C', 'D'], size=rows_per_sheet)
            }
            df = pd.DataFrame(data)
            df.to_excel(writer, sheet_name=sheet_name, index=False)
            
            worksheet = writer.sheets[sheet_name]
            
            # Add a native Excel chart (acting as flowchart/visual representation)
            chart = workbook.add_chart({'type': 'column'})
            chart.add_series({
                'name':       f'={sheet_name}!$D$1',
                'categories': f'={sheet_name}!$A$2:$A$100',
                'values':     f'={sheet_name}!$D$2:$D$100',
            })
            chart.set_title({'name': 'Salary Overview'})
            worksheet.insert_chart('J2', chart)
            
            # Add an image in a few sheets
            if i % 3 == 0:
                worksheet.insert_image('J20', 'dummy_chart.png')
                
    print(f"Created {excel_file} with {num_sheets} sheets, charts, and images.")

# 2. Generate Word Document (Lots of data)
def create_docx():
    doc = Document()
    doc.add_heading('Comprehensive Corporate Handbook & Reports', 0)
    
    for section in range(1, 15):
        doc.add_heading(f'Section {section}: {fake.bs().title()}', level=1)
        doc.add_paragraph(' '.join(fake.paragraphs(nb=10)))
        
        # Add a big table every few sections
        if section % 3 == 0:
            doc.add_heading('Data Table Summary', level=2)
            table = doc.add_table(rows=1, cols=4)
            hdr_cells = table.rows[0].cells
            hdr_cells[0].text = 'Id'
            hdr_cells[1].text = 'Name'
            hdr_cells[2].text = 'Region'
            hdr_cells[3].text = 'Sales'
            
            for _ in range(80):
                row_cells = table.add_row().cells
                row_cells[0].text = str(random.randint(1000, 9999))
                row_cells[1].text = fake.name()
                row_cells[2].text = fake.country()
                row_cells[3].text = f"${random.randint(100, 10000)}"
                
        doc.add_page_break()
            
    doc.save('sample_complex.docx')
    print("Created sample_complex.docx with large amount of text and tables.")

# 3. Generate Markdown
def create_md():
    md_content = """# Architecture Guidelines
## 1. Microservices
All new services must be written in Python 3.11+ or Go. 
### 1.1 Communication
Services must communicate asynchronously via Kafka.
## 2. Database
PostgreSQL is the default relational database. NoSQL use cases should default to MongoDB.
"""
    with open('sample.md', 'w') as f:
        f.write(md_content)
    print("Created sample.md")

# 4. Generate OpenAPI YAML
def create_yaml():
    yaml_content = {
        "openapi": "3.0.0",
        "info": {"title": "User API", "version": "1.0.0"},
        "paths": {
            "/users": {
                "get": {
                    "summary": "Get all users",
                    "responses": {"200": {"description": "List of users"}}
                },
                "post": {
                    "summary": "Create user",
                    "requestBody": {
                        "content": {"application/json": {"schema": {"type": "object", "properties": {"name": {"type": "string"}}}}}
                    },
                    "responses": {"201": {"description": "User created"}}
                }
            }
        }
    }
    with open('sample_openapi.yaml', 'w') as f:
        yaml.dump(yaml_content, f, sort_keys=False)
    print("Created sample_openapi.yaml")

# 5. Generate Swagger JSON
def create_json():
    json_content = {
        "swagger": "2.0",
        "info": {"title": "Payment API", "version": "1.0"},
        "paths": {
            "/pay": {
                "post": {
                    "summary": "Process payment",
                    "parameters": [{"name": "amount", "in": "body", "required": True}],
                    "responses": {"200": {"description": "Payment successful"}}
                }
            }
        }
    }
    with open('sample_swagger.json', 'w') as f:
        json.dump(json_content, f, indent=2)
    print("Created sample_swagger.json")

if __name__ == "__main__":
    print("Generating complicated sample files...")
    create_excel()
    create_docx()
    create_md()
    create_yaml()
    create_json()
    print("Done!")